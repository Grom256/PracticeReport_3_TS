# -*- coding: utf-8 -*-
"""
Generate a concise, concentrated version of Section 2 for the GromBricks Practice Report.
"""
import os
import io
import docx
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_heading_1(doc, text):
    p = doc.add_paragraph(style='Heading 1')
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.first_line_indent = Inches(0)
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(14)
    run.font.bold = True
    return p

def add_heading_2(doc, text):
    p = doc.add_paragraph(style='Heading 2')
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.first_line_indent = Inches(0.492)
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(14)
    run.font.bold = True
    return p

def add_subheading(doc, text):
    p = doc.add_paragraph(style='Normal')
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.first_line_indent = Inches(0.492)
    p.paragraph_format.line_spacing = 1.5
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(14)
    run.font.bold = True
    return p

def add_paragraph(doc, text):
    p = doc.add_paragraph(style='Normal')
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.first_line_indent = Inches(0.492)
    p.paragraph_format.line_spacing = 1.5
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(14)
    return p

def add_placeholder(doc, text):
    p = doc.add_paragraph(style='Normal')
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.first_line_indent = Inches(0)
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.italic = True
    return p

def add_table_title(doc, text):
    p = doc.add_paragraph(style='Normal')
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.first_line_indent = Inches(0)
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    return p

def add_table(doc, headers, data, col_widths=None):
    table = doc.add_table(rows=len(data) + 1, cols=len(headers), style='Table Grid')
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    hdr_cells = table.rows[0].cells
    for i, header_text in enumerate(headers):
        hdr_cells[i].text = header_text
        set_cell_margins(hdr_cells[i], top=100, bottom=100, left=130, right=130)
        p = hdr_cells[i].paragraphs[0]
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Inches(0)
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        for r in p.runs:
            r.font.name = 'Times New Roman'
            r.font.size = Pt(11)
            r.font.bold = True

    for row_idx, row_data in enumerate(data):
        row_cells = table.rows[row_idx + 1].cells
        for col_idx, cell_value in enumerate(row_data):
            row_cells[col_idx].text = cell_value
            set_cell_margins(row_cells[col_idx], top=80, bottom=80, left=130, right=130)
            p = row_cells[col_idx].paragraphs[0]
            p.paragraph_format.first_line_indent = Inches(0)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            if col_idx == 0:
                p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            else:
                p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            for r in p.runs:
                r.font.name = 'Times New Roman'
                r.font.size = Pt(11)

    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Inches(w)

    doc.add_paragraph(style='Normal').paragraph_format.space_after = Pt(3)
    return table

def build_condensed_section_2(doc):
    # =========================================================================
    # РОЗДІЛ 2
    # =========================================================================
    add_heading_1(doc, "РОЗДІЛ 2. РОЗРОБКА FULLSTACK-ЗАСТОСУНКУ НА PAYLOAD CMS ТА TYPESCRIPT (ГУСАРОВ, ЮРЧИШИН)")

    # -------------------------------------------------------------------------
    # 2.1. Архітектура та Payload CMS
    # -------------------------------------------------------------------------
    add_heading_2(doc, "2.1. Архітектура та Payload CMS: монорепозиторій, конфігурація Payload (колекції, типи блоків, медіа), хмарна БД та кастомні API-ендпоінти.")

    add_paragraph(doc, "Другий етап навчальної практики присвячено розробці Fullstack-застосунку інтернет-магазину GromBricks на основі безголової Payload CMS 3.x, платформи Next.js, мови TypeScript, збирача Vite та хмарної реляційної СУБД Neon PostgreSQL. На відміну від монолітної архітектури Django MVT, тут реалізовано decoupled-підхід із розділенням відповідальності між Headless CMS та автономною вітриною магазину.")

    add_paragraph(doc, "Для структурування кодової бази використано архітектуру монорепозиторію з механізмом npm workspaces. Вона об'єднує підсистему бекенду (папка cms/) та фронтенду (папка frontend/) під єдиним контролем версій Git. Кореневий файл package.json оголошує спільні скрипти життєвого циклу: dev:cms (запуск Payload CMS на порту 3000), dev:frontend (запуск Vite на порту 5173) та build (конвеєрна продуктова компіляція обох частин).")

    add_placeholder(doc, "[ВСТАВИТИ ЗОБРАЖЕННЯ ЛІСТИНГУ 2.1: Конфігурація монорепозиторію та робочих просторів package.json]")

    add_paragraph(doc, "Критично важливим рішенням стала відмова від локальних файлових баз даних (SQLite) на користь керованої хмарної СУБД Neon PostgreSQL 16. Цільова платформа розгортання Vercel функціонує за безсерверною моделлю (Serverless Functions) з ефемерною файловою системою, тому будь-які локальні файли безповоротно стираються після перезапуску контейнера. Хмарна PostgreSQL у Neon підтримує постійне сховище, пулінг з'єднань через PgBouncer та захищене з'єднання SSL (ssl: { rejectUnauthorized: false }), що налаштовано в адаптері @payloadcms/db-postgres.")

    add_paragraph(doc, "Головний конфігураційний файл cms/src/payload.config.ts реєструє точку входу в адмін-панель (/admin), автентифікацію Users.slug, редактор Lexical (@payloadcms/richtext-lexical), модуль Sharp для оптимізації медіа та генерацію типів payload-types.ts. Для безперешкодної взаємодії вітрини з бекендом як локально, так і у продуктовому середовищі Vercel, активовано політику cors: '*' та вимкнено блокування csrf.")

    add_placeholder(doc, "[ВСТАВИТИ ЗОБРАЖЕННЯ ЛІСТИНГУ 2.2: Конфігурація Payload CMS та підключення хмарної бази даних (payload.config.ts)]")

    add_paragraph(doc, "Модель даних побудована на декларативних колекціях (Collections) та глобальних об'єктах (Globals):")
    add_paragraph(doc, "• Products — каталог товарів із полями: name, slug, унікальний артикул sku, тип productType (minifigure/part), серія series, ціна price, стан condition (новий/вживаний), залишок stock, прапорець рідкісності isRare, реляція category, зображення imageUrl та статус активності isActive;\n"
                       "• Categories — рубрикатор асортименту (title, slug, description, badge, iconName);\n"
                       "• Orders — замовлення клієнтів із відкритим публічним доступом на створення: orderNumber (формату GB-2026-XXXX), контактні дані клієнта, служба доставки deliveryService (Нова Пошта, Укрпошта, Meest), адреса, спосіб оплати, статус status та масив товарних позицій items із зафіксованою ціною unitPrice;\n"
                       "• Pages — сторінки з динамічним блоковим конструктором layout (Page Builder);\n"
                       "• Media — збереження зображень з автоматичною генерацією розмірів Sharp (thumbnail 300×300, card 600×600, hero 1440×600);\n"
                       "• Users — системні користувачі з підтримкою JWT-автентифікації та ролями admin, editor, customer;\n"
                       "• HeaderSettings та FooterSettings (Globals) — наскрізні параметри навігації, телефони, посилання на Telegram та копірайт.")

    add_table_title(doc, "Таблиця 2.1 - Основні колекції та глобальні налаштування Payload CMS")
    add_table(doc,
              ["Сутність / Slug", "Тип у Payload", "Призначення та основні поля", "Правила доступу"],
              [
                  ["products", "Collection", "Каталог деталей та фігурок: name, sku, productType, series, price, stock, isRare, category, imageUrl, isActive.", "read: public, write: admin/editor"],
                  ["categories", "Collection", "Рубрикатор асортименту: title, slug, description, badge, iconName.", "read: public, write: admin"],
                  ["orders", "Collection", "Замовлення: orderNumber, customerName, customerPhone, deliveryService, status, totalAmount, items[].", "read: public, create: public, update: admin"],
                  ["pages", "Collection", "Сторінки з блоковим конструктором: title, slug, isPublished, layout (масив blocks).", "read: public, write: admin/editor"],
                  ["media", "Collection", "Завантаження та оптимізація зображень Sharp: alt, imageSizes (thumbnail, card, hero).", "read: public, upload: admin/editor"],
                  ["users", "Collection", "Системні користувачі: name, email, role (admin / editor / customer), phone.", "auth: true, read: public, manage: admin"],
                  ["header / footer", "Global", "Наскрізні параметри шапки та підвалу: контакти, соцмережі, навігація, копірайт.", "read: public, write: admin"]
              ],
              col_widths=[1.4, 1.1, 3.2, 1.3])

    add_placeholder(doc, "[ВСТАВИТИ ЗОБРАЖЕННЯ ЛІСТИНГУ 2.3: Схема колекції Products та визначення полів (Products.ts)]")

    add_paragraph(doc, "Блоковий конструктор (директорія cms/src/blocks/) містить ізольовані модулі: HeroBlock (головний банер із кнопками), FeaturesBlock (переваги магазину з іконками), ProductGridBlock (товарна сітка з правилами фільтрації), BannerBlock (промо-акції) та CallToActionBlock (блок зворотного зв'язку через Telegram).")

    add_paragraph(doc, "Стандартний REST API Payload повертає надлишково вкладені JSON-відповіді з метаданими. Для оптимізації мережевого навантаження та спрощення парсингу на чистому клієнті реалізовано патерн Backend for Frontend (BFF) через кастомні API-ендпоінти у cms/src/endpoints/:")
    add_paragraph(doc, "• GET /api/common — віддає плоский об'єкт глобальних налаштувань із вбудованим fallback-механізмом;\n"
                       "• GET /api/page?path=/ — повертає готову структуру блоків сторінки, автоматично наповнюючи блок productGrid актуальними товарами за один HTTP-запит;\n"
                       "• GET /api/products/search — виконує повнотекстовий пошук (q) та фільтрацію каталогу за типом, серією, рідкісністю й ціною;\n"
                       "• POST /api/cart/validate — валідує складські залишки та ціни товарів у кошику;\n"
                       "• POST /api/orders/checkout — перевіряє реквізити, генерує унікальний номер замовлення (GB-2026-XXXX), створює запис у колекції orders та атомарно зменшує залишок stock у колекції products;\n"
                       "• POST /api/seed — автоматизовано заповнює базу 5 базовими категоріями, 8 товарами та головною сторінкою.")

    add_table_title(doc, "Таблиця 2.2 - Кастомні API-ендпоінти сервера GromBricks")
    add_table(doc,
              ["HTTP-метод та URL", "Модуль обробника", "Вхідні параметри", "Формат та опис відповіді"],
              [
                  ["GET /api/common", "commonEndpoint.ts", "Відсутні", "{ success: true, data: { header, footer } } — глобальні налаштування."],
                  ["GET /api/page", "pageEndpoint.ts", "path (шлях сторінки, наприклад /)", "{ success: true, data: { title, slug, blocks[] } } — плоскі блоки."],
                  ["GET /api/products/search", "searchEndpoint.ts", "q, type, series, condition, onlyRare, sortBy", "{ success: true, data: Product[] } — відфільтрований каталог."],
                  ["POST /api/cart/validate", "cartEndpoints.ts", "JSON: { items: [{ productId, quantity }] }", "{ success: true, data: CartValidationResult } — перевірка наявності."],
                  ["POST /api/orders/checkout", "orderEndpoints.ts", "JSON: { customer, delivery, paymentMethod, items }", "{ success: true, data: { orderId, orderNumber, totalAmount } } — списання залишку."],
                  ["POST /api/seed", "seedEndpoint.ts", "Відсутні", "{ success: true, message, categoriesCount, productsCount } — наповнення БД."]
              ],
              col_widths=[1.7, 1.4, 1.8, 2.1])

    add_placeholder(doc, "[ВСТАВИТИ ЗОБРАЖЕННЯ ЛІСТИНГУ 2.4: Транзакційне оформлення замовлення у файлі orderEndpoints.ts]")

    # -------------------------------------------------------------------------
    # 2.2. Клієнтська частина та стани UI
    # -------------------------------------------------------------------------
    add_heading_2(doc, "2.2. Клієнтська частина та стани UI: асинхронні запити до API, шаблонізація блоків, адаптивна верстка, обробка станів (завантаження / помилка / готовність).")

    add_paragraph(doc, "Клієнтська вітрина розроблена у директорії frontend/ відповідно до вимог практики: без важких фреймворків (React, Angular чи Vue), на чистому Vanilla TypeScript, семантичному HTML5, модульному CSS3 та збирачі Vite. Це забезпечує мінімальний обсяг бандла, нульовий оверхед віртуального DOM та швидкий рендеринг.")

    add_paragraph(doc, "Застосунок реалізовано за багатосторінковою моделлю (MPA): index.html (презентаційний лендінг із 3D-вітриною), catalog.html (каталог із живим пошуком та фільтрами), cart.html (кошик із синхронізацією у localStorage) та checkout.html (форма замовлення з валідацією).")

    add_paragraph(doc, "Мережевий шар винесено в модуль frontend/src/api.ts. Запити виконуються через нативний fetch (async/await) за базовою адресою VITE_API_BASE. Для запобігання збоям інтерфейсу всі виклики захищені блоками try...catch із поверненням структурованих резервних даних (fallback).")

    add_placeholder(doc, "[ВСТАВИТИ ЗОБРАЖЕННЯ ЛІСТИНГУ 2.5: Асинхронні запити до API та механізм fallback-даних (api.ts)]")

    add_paragraph(doc, "Рендеринг контенту організовано за патерном «Фабрика шаблонів» у frontend/src/blocks/renderBlocks.ts. Функція renderBlocks ітерує масив блоків і через switch (blockType) делегує генерацію розмітки функціям renderHeroBlock, renderFeaturesBlock, renderCategoryCardsBlock, renderProductGridBlock та renderCtaBlock. Додатково модуль modelViewer.ts ініціалізує тривимірну сцену Three.js із динамічним освітленням, орбітальним керуванням OrbitControls та підтримкою моделей LEGO (Kai, Cole) у форматах OBJ/MTL та DAE.")

    add_placeholder(doc, "[ВСТАВИТИ ЗОБРАЖЕННЯ ЛІСТИНГУ 2.6: Реалізація шаблонізатора сторінки в renderBlocks.ts]")

    add_paragraph(doc, "Адаптивна верстка (frontend/src/style.css) реалізована на CSS Grid (repeat(auto-fill, minmax(280px, 1fr))) та Flexbox у фірмовій темній палітрі (#05070d, #070b16) із золотими акцентами (#e6bf55, #ffd700). Завдяки медіа-запитам (1024px, 768px, 480px) на смартфонах активується мобільна шторка навігації (.mobile-nav-drawer), товарна сітка перебудовується в 1 колонку, а таблиця кошика трансформується у вертикальні картки.")

    add_paragraph(doc, "Життєвий цикл інтерфейсу обробляє 4 канонічні стани UI (UI States):")
    add_paragraph(doc, "1. Loading — показ анімованого індикатора .lego-spinner та тексту очікування під час виконання запиту;\n"
                       "2. Success — відображення згенерованої структури блоків або сітки товарів із підключенням слухачів подій;\n"
                       "3. Empty — відображення інформативного блоку .empty-container із кнопкою повернення до каталогу при нульових результатах пошуку або порожньому кошику;\n"
                       "4. Error — виведення банера помилки .checkout-error-banner та блокування повторної відправки форми.")

    add_table_title(doc, "Таблиця 2.3 - Стани користувацького інтерфейсу вітрини магазину та їх обробка")
    add_table(doc,
              ["Стан інтерфейсу", "Умова виникнення", "Візуальний компонент та CSS-клас", "Поведінка системи та зворотний зв'язок"],
              [
                  ["Loading (Завантаження)", "Початок виконання fetch-запиту до /api/page або /api/products/search.", ".loading-container, .lego-spinner, .loading-text", "Анімація обертання кубика LEGO, блокування кліків."],
                  ["Success (Готовність)", "Отримання коду 200/201 та валідного масиву даних.", ".catalog-grid, .cart-page-grid, .confirmation-card", "Рендеринг товарів або блоків, ініціалізація слухачів подій."],
                  ["Empty (Порожньо)", "0 знайдених товарів або 0 позицій у кошику.", ".empty-container, .empty-title, .empty-message", "Повідомлення з кнопкою переходу до повного каталогу."],
                  ["Error (Помилка)", "Мережевий збій або помилка валідації 400 Bad Request.", ".error-container, .checkout-error-banner", "Інформативне повідомлення, розблокування полів для виправлення."]
              ],
              col_widths=[1.5, 1.8, 1.8, 1.9])

    # -------------------------------------------------------------------------
    # 2.3. Типізація на TypeScript
    # -------------------------------------------------------------------------
    add_heading_2(doc, "2.3. Типізація на TypeScript: створення типів для структури сторінки та її блоків, використання Type Aliases, Union Types та Intersection Types для опису різних типів даних, типізація функцій і об'єктів, типізація даних, отриманих через API.")

    add_paragraph(doc, "Статична типізація TypeScript забезпечує передбачуваність та надійність Fullstack-застосунку GromBricks. Вона усуває помилки часу виконання та визначає строгі контракти між базою даних, бекендом і фронтендом (файли cms/src/types/shop.ts та frontend/src/types.ts).")

    add_paragraph(doc, "Псевдоніми типів (Type Aliases) замінюють небезпечні «магічні рядки»: ID = string | number, ProductType = 'minifigure' | 'part', ProductCondition = 'new' | 'used', OrderStatus = 'pending' | 'processing' | 'shipped' | 'delivered' | 'cancelled', DeliveryService = 'nova_poshta' | 'ukrposhta' | 'meest', PaymentMethod, LegoSeries, Currency, SortOption.")

    add_paragraph(doc, "Типи об'єднання (Union Types) та дискриміновані об'єднання (Discriminated Unions) застосовано для поліморфних моделей:")
    add_paragraph(doc, "1. Дискриміноване об'єднання блоків сторінки: інтерфейси HeroBlockData, ProductGridBlockData, FeaturesBlockData, BannerBlockData, CtaBlockData, ModelViewerBlockData містять обов'язкове спільне поле blockType. Їх об'єднання формує тип PageBlockData. Це дає змогу компілятору виконувати звуження типу (type narrowing) у конструкціях switch (block.blockType).\n"
                       "2. Дискримінований стан UI: export type UIState<T> = { status: 'idle' } | { status: 'loading' } | { status: 'success'; data: T } | { status: 'empty' } | { status: 'error'; error: string }.\n"
                       "3. Уніфікована відповідь сервера: export type ApiResponse<T> = { success: true; data: T; timestamp: string } | { success: false; error: string; timestamp: string }.")

    add_placeholder(doc, "[ВСТАВИТИ ЗОБРАЖЕННЯ ЛІСТИНГУ 2.7: Система типів: Discriminated Unions та стани UI (cms/src/types/shop.ts)]")

    add_paragraph(doc, "Типи перетину (Intersection Types, оператор &) використано для розширення сутностей без дублювання коду:")
    add_paragraph(doc, "• ProductWithCategory = Product & { categoryDetails: Category } — гарантує наявність об'єкта категорії;\n"
                       "• CartItem = { productId: ID; quantity: number } & Pick<Product, 'name' | 'slug' | 'sku' | 'price' | 'imageUrl' | 'stock'> — комбінує ідентифікатор із безпечним знімком товару;\n"
                       "• CartItemWithDetails = CartItem & { product: Product; subtotal: number; isAvailable: boolean; hasEnoughStock: boolean } — містить повні розрахунки для інтерфейсу;\n"
                       "• EnrichedOrder = Order & { formattedDate: string; itemsCount: number; isShipped: boolean }.")

    add_paragraph(doc, "У модулі cms/src/lib/shopUtils.ts типізовано сигнатури функцій та об'єктів:")
    add_paragraph(doc, "• CartValidatorFn — функція validateCart(items, productsPool) звіряє запитану кількість із залишком product.stock та повертає типізований результат CartValidationResult;\n"
                       "• PriceCalculatorFn — функція calculateOrderPricing(items, deliveryService) обчислює суму замовлення та вартість доставки (безкоштовно від 1500 грн);\n"
                       "• FilterProductsFn — функція filterProducts(products, options) виконує багатофакторний відбір за комплексним об'єктом опцій ProductFilterOptions.")

    add_placeholder(doc, "[ВСТАВИТИ ЗОБРАЖЕННЯ ЛІСТИНГУ 2.8: Типізовані сервісні функції магазину (cms/src/lib/shopUtils.ts)]")

    add_paragraph(doc, "Типізація даних API поєднує автоматично згенерований файл payload-types.ts (схеми бази даних) та інтерфейси передачі даних (DTO): CheckoutApiRequest (контракт відправки замовлення з фронтенду) і CheckoutApiResponseData (контракт підтвердження замовлення).")

    # -------------------------------------------------------------------------
    # 2.4. Деплой та контрольний приклад
    # -------------------------------------------------------------------------
    add_heading_2(doc, "2.4. Деплой та контрольний приклад: розгортання проєкту, оформлення репозиторію (README.md), контрольний приклад із скріншотами (редактор в CMS, відвідувач на сайті, обробка помилок).")

    add_paragraph(doc, "Промислове розгортання проєкту реалізовано у тандемі хмарних платформ Neon та Vercel. У Neon створено базу даних PostgreSQL 16 із параметрами захищеного підключення sslmode=require та пулером PgBouncer. У панелі Vercel із єдиного репозиторію створено два проєкти: grombricks-cms (корінь cms, Next.js, змінні DATABASE_URL, PAYLOAD_SECRET, NEXT_PUBLIC_SERVER_URL) та grombricks-front (корінь frontend, вихідна папка dist, змінна VITE_API_BASE). Це повністю нівелює проблему ефемерної файлової системи Vercel.")

    add_table_title(doc, "Таблиця 2.4 - Конфігураційні змінні оточення та параметри розгортання")
    add_table(doc,
              ["Застосунок у Vercel", "Параметр / Змінна оточення", "Значення / Конфігурація", "Призначення"],
              [
                  ["grombricks-cms (Backend)", "Root Directory", "cms", "Вказує Vercel на кореневу папку сервера Next.js."],
                  ["grombricks-cms (Backend)", "DATABASE_URL", "postgres://user:pass@ep-xyz.neon.tech/grombricks_db?sslmode=require", "Підключення до хмарної бази Neon PostgreSQL."],
                  ["grombricks-cms (Backend)", "PAYLOAD_SECRET", "32-символьний криптографічний хеш", "Захист JWT-сесій користувачів CMS."],
                  ["grombricks-cms (Backend)", "NEXT_PUBLIC_SERVER_URL", "https://grombricks-cms.vercel.app", "Публічний домен сервера CMS та API."],
                  ["grombricks-front (Frontend)", "Root Directory / Output", "frontend / dist", "Шлях до вітрини Vite та каталогу продуктової збірки."],
                  ["grombricks-front (Frontend)", "VITE_API_BASE", "https://grombricks-cms.vercel.app", "Базовий URL для cross-origin fetch-запитів до API."]
              ],
              col_widths=[1.6, 1.8, 2.0, 1.6])

    add_paragraph(doc, "Документація README.md у корені репозиторію містить опис структури монорепозиторію, системні вимоги (Node.js 20.18+, Neon), інструкцію із встановлення залежностей (npm install), перелік npm-скриптів, зведення API-ендпоінтів та кроки розгортання у хмарі Vercel.")

    add_paragraph(doc, "Працездатність системи підтверджено тестуванням трьох сценаріїв контрольного прикладу:")

    add_subheading(doc, "Сценарій 1. Робота адміністратора та редактора контенту в системі Payload CMS")
    add_paragraph(doc, "1. Адміністратор авторизується в адмін-панелі за адресою /admin. Відображається перелік колекцій: Products, Categories, Orders, Pages, Media, Users.")
    add_placeholder(doc, "[ВСТАВИТИ РИСУНОК 2.2: Панель керування Payload CMS: список товарів колекції Products]")

    add_paragraph(doc, "2. У колекції Products відкривається картка товару «Дарт Вейдер», встановлюється прапорець isRare=true та оновлюється залишок stock з 5 до 10 одиниць. Зміни миттєво фіксуються в Neon PostgreSQL.")
    add_paragraph(doc, "3. У колекції Pages для сторінки «home» змінюється заголовок блоку Hero на «Ексклюзивні мініфігурки LEGO 2026», переставляються блоки у Page Builder та зберігаються оновлення.")
    add_placeholder(doc, "[ВСТАВИТИ РИСУНОК 2.3: Візуальний конструктор блоків сторінки (Page Builder) у CMS]")

    add_paragraph(doc, "4. У Globals коригуються контакти шапки HeaderSettings, а в колекції Orders нове замовлення клієнта переводиться зі статусу pending у processing.")
    add_placeholder(doc, "[ВСТАВИТИ РИСУНОК 2.4: Детальний перегляд замовлення та зміна статусу в колекції Orders]")

    add_subheading(doc, "Сценарій 2. Сценарій взаємодії покупця з вітриною інтернет-магазину")
    add_paragraph(doc, "1. Відвідувач відкриває index.html. Асинхронні запити завантажують налаштування та блоки. Шаблонізатор відображає Hero-екран, блок переваг та ініціалізує 3D-сцену Three.js із можливістю обертання фігурки ніндзя.")
    add_placeholder(doc, "[ВСТАВИТИ РИСУНОК 2.5: Головна сторінка магазину GromBricks з інтерактивною 3D-вітриною]")

    add_paragraph(doc, "2. На сторінці catalog.html користувач шукає «Rex» і позначає чекбокс «Рідкісні ексклюзиви». Застосунок через /api/products/search?q=Rex&onlyRare=true знаходить фігурку Клон-капітана Рекса (1450 грн).")
    add_placeholder(doc, "[ВСТАВИТИ РИСУНОК 2.6: Каталог товарів із живим пошуком та фільтрацією за параметрами]")

    add_paragraph(doc, "3. Після натискання «В кошик» товар додається в localStorage, а лічильник у шапці змінюється з 0 на 1. На сторінці cart.html кількість збільшується до 2 шт., а сума автоматично перераховується до 2900 грн.")
    add_placeholder(doc, "[ВСТАВИТИ РИСУНОК 2.7: Сторінка кошика покупця з розрахунком вартості]")

    add_paragraph(doc, "4. На сторінці checkout.html клієнт заповнює реквізити («Олександр Мельник», місто Львів, Нова Пошта №15). Після відправки сервер через /api/orders/checkout списує 2 одиниці зі складу в БД та повертає номер замовлення GB-2026-4819. Кошик очищується.")
    add_placeholder(doc, "[ВСТАВИТИ РИСУНОК 2.8: Підтвердження успішно створеного замовлення]")

    add_subheading(doc, "Сценарій 3. Обробка граничних станів та помилок (Error Handling)")
    add_paragraph(doc, "1. Під час запитів до API відображається плавний анімаційний спінер .lego-spinner, що запобігає повторним клікам.")
    add_paragraph(doc, "2. Запит за неіснуючим артикулом коректно виводить блок .empty-container із повідомленням «Нічого не знайдено».")
    add_placeholder(doc, "[ВСТАВИТИ РИСУНОК 2.9: Стан інтерфейсу при відсутності результатів пошуку в каталозі]")

    add_paragraph(doc, "3. При спробі замовити кількість товару, яка перевищує наявний залишок stock у базі даних (наприклад, 100 одиниць шолома при залишку 10), ендпоінт checkout відхиляє операцію (400 Bad Request), а вітрина виводить червоний банер .checkout-error-banner.")
    add_placeholder(doc, "[ВСТАВИТИ РИСУНОК 2.10: Повідомлення про помилку валідації залишків на складі при оформленні]")

    add_paragraph(doc, "4. У разі недоступності сервера бекенду модуль api.ts автоматично перемикається на вбудовані fallback-структури, забезпечуючи стабільну роботу вітрини.")

    # -------------------------------------------------------------------------
    # Висновок до розділу 2
    # -------------------------------------------------------------------------
    add_subheading(doc, "Висновок до розділу 2")
    add_paragraph(doc, "У другому розділі звіту спроєктовано та реалізовано Fullstack-застосунок GromBricks за архітектурою монорепозиторію з використанням Payload CMS 3.x, хмарної СУБД Neon PostgreSQL, платформи Vite та строгої типізації TypeScript. Розроблено 6 кастомних API-ендпоінтів для оптимізації передачі даних і транзакційного контролю складських залишків. Проблему ефемерної файлової системи безсерверного середовища Vercel вирішено через зовнішню реляційну базу даних Neon. Клієнтська вітрина побудована за стандартом MPA на чистому TypeScript без сторонніх UI-фреймворків, із реалізацією адаптивної верстки, шаблонізатора блоків, інтерактивної 3D-вітрини Three.js та надійної обробки життєвих циклів UI. Контрольний приклад підтвердив повну працездатність усіх компонентів системи.")

if __name__ == '__main__':
    src_user_upload = r'C:\Users\vitai\.gemini\antigravity\brain\85df3c14-278b-43f5-ae1e-2ffdec46da01\.user_uploaded\media_1793085605805_620abf0d.docx'
    dst_root = r'd:\PracticeReport_3_TS\Звіт_з_практики_Розділ_1_і_2.docx'
    dst_docs = r'd:\PracticeReport_3_TS\docs\Звіт_з_практики_Розділ_1_і_2.docx'

    with open(src_user_upload, 'rb') as f:
        buf = io.BytesIO(f.read())

    doc = docx.Document(buf)
    body = doc._body._body
    elements = list(body)
    
    # Paragraph 152 is the last paragraph of Section 1
    p152 = doc.paragraphs[152]._p
    idx_p152 = elements.index(p152)

    # Remove all elements after p152 except sectPr to start clean
    for el in elements[idx_p152 + 1:]:
        if not el.tag.endswith('sectPr'):
            body.remove(el)

    print(f"Cleaned up to Section 1: {len(doc.paragraphs)} paragraphs remain.")

    # Build the condensed Section 2
    build_condensed_section_2(doc)
    print(f"Total paragraphs now: {len(doc.paragraphs)}")
    print(f"Total tables now: {len(doc.tables)}")

    # Save to docs folder
    os.makedirs(os.path.dirname(dst_docs), exist_ok=True)
    doc.save(dst_docs)
    print(f"Saved successfully to: {dst_docs}")

    # Try to save to root if not locked, otherwise save as _скорочений.docx
    try:
        doc.save(dst_root)
        print(f"Saved successfully to: {dst_root}")
    except PermissionError:
        alt_root = r'd:\PracticeReport_3_TS\Звіт_з_практики_Розділ_1_і_2_скорочений.docx'
        doc.save(alt_root)
        print(f"Original file was open in Word, saved condensed copy to: {alt_root}")

    print("All tasks completed successfully!")
