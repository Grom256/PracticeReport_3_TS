# -*- coding: utf-8 -*-
"""
Повне оновлення Розділу 2 відповідно до реального стану кодової бази GromBricks:
- Повне динамічне редагування всіх текстів сторінки через Payload CMS
- Новий текстовий блок TextBlock для створення базових інформаційних сторінок (Про нас, Доставка)
- Клієнтська SPA-маршрутизація (Router) на чистому Vanilla TypeScript з хлібними крихтами та 404
- Оновлений бейдж Hero (ПРЕМІАЛЬНІ LEGO КОЛЕКЦІЇ без емодзі), зменшений адаптивний розмір заголовка
- Якірна навігація для кнопки «ПЕРЕВАГИ» з іконкою щита без виринаючих вікон магазину
- Інтерактивна 3D-вітрина Three.js, повна типізація TypeScript, хмарний деплой Vercel + Neon
- Форматування згідно з вимогами: Times New Roman 14pt, міжрядковий 1.5, абзацний відступ 1.25 см
"""
import os
import copy
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_margins(cell, top=100, bottom=100, left=130, right=130):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_heading_1(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.line_spacing = 1.5
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(14)
    run.font.bold = True
    return p

def add_heading_2(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Pt(35.45) # ~1.25 cm
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.5
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(14)
    run.font.bold = True
    return p

def add_paragraph(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Pt(35.45)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.5
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(14)
    return p

def add_bullet(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Pt(35.45)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.5
    run = p.add_run(f"• {text}")
    run.font.name = 'Times New Roman'
    run.font.size = Pt(14)
    return p

def add_placeholder(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.5
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    run.font.italic = True
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 51, 102)
    return p

def add_table_title(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    run.font.bold = True
    return p

def add_table(doc, headers, data, col_widths=None):
    table = doc.add_table(rows=len(data) + 1, cols=len(headers), style='Table Grid')
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    hdr_cells = table.rows[0].cells
    for i, header_text in enumerate(headers):
        hdr_cells[i].text = header_text
        set_cell_margins(hdr_cells[i], top=100, bottom=100, left=120, right=120)
        p = hdr_cells[i].paragraphs[0]
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Pt(0)
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        for r in p.runs:
            r.font.name = 'Times New Roman'
            r.font.size = Pt(11)
            r.font.bold = True

    for row_idx, row_data in enumerate(data):
        row_cells = table.rows[row_idx + 1].cells
        for col_idx, cell_value in enumerate(row_data):
            row_cells[col_idx].text = cell_value
            set_cell_margins(row_cells[col_idx], top=80, bottom=80, left=120, right=120)
            p = row_cells[col_idx].paragraphs[0]
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            if col_idx == 0:
                p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            else:
                p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            for r in p.runs:
                r.font.name = 'Times New Roman'
                r.font.size = Pt(11)

    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Inches(w)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return table

def generate_clean_landing_section_2(doc):
    # =========================================================================
    # РОЗДІЛ 2
    # =========================================================================
    add_heading_1(doc, "РОЗДІЛ 2. Розробка Fullstack-застосунку на Payload CMS та TypeScript")

    # -------------------------------------------------------------------------
    # 2.1. Архітектура та Payload CMS
    # -------------------------------------------------------------------------
    add_heading_2(doc, "2.1. Архітектура та Payload CMS: монорепозиторій, конфігурація Payload (колекції, типи блоків, медіа), хмарна БД та кастомні API-ендпоінти.")

    add_paragraph(doc, "Під час виконання другого етапу виробничої практики основну увагу було зосереджено на розробці сучасного інтерактивного презентаційного лендінгу (Landing Page) для бренду GromBricks, що спеціалізується на колекційних фігурках і деталях LEGO. На відміну від монолітної структури першого розділу, архітектура цього етапу побудована за decoupled-моделлю (розділена Fullstack-система). Вона передбачає суворе розділення системи на сервер керування контентом (Payload CMS 3.x) та легку, швидку клієнтську вітрину на чистому Vanilla TypeScript під керуванням збирача Vite.")

    add_paragraph(doc, "Для забезпечення злагодженої роботи двох підсистем у межах єдиного середовища кодову базу організовано як монорепозиторій за допомогою менеджера пакетів npm та технології npm workspaces. Початкову структуру було створено командою npm init, а загальний маніфест package.json у корені містить метадані проєкту, спільні налаштування та конфігурацію робочих просторів:")

    add_bullet(doc, "директорія cms/ — серверна частина на базі Payload CMS 3.x та Next.js, що відповідає за зберігання контенту, автентифікацію адміністраторів, візуальний блоковий конструктор Page Builder та медіафайли;")
    add_bullet(doc, "директорія frontend/ — клієнтська презентаційна вітрина на чистому Vanilla TypeScript, HTML5 та CSS3 під керуванням збирача Vite з інтегрованою 3D-вітриною Three.js та легким клієнтським маршрутизатором;")
    add_bullet(doc, "розподіл залежностей — у package.json сторонні бібліотеки чітко розділено на dependencies (пакети, необхідні для безпосередньої роботи додатка у продакшені: payload, @payloadcms/next, @payloadcms/db-postgres, three, sharp) та devDependencies (інструменти компіляції TypeScript, збірки Vite, тайпінги та лінтери, потрібні виключно на етапі розробки);")
    add_bullet(doc, "спільні скрипти життєвого циклу — у секції scripts налаштовано зручні аліаси: npm run dev:cms (запуск адмін-панелі на порті 3000), npm run dev:frontend (запуск клієнтського середовища на порті 5173) та npm run build для компіляції обох частин;")
    add_bullet(doc, "ігнорування службових каталогів — директорію node_modules внесено до файлу .gitignore через її великий обсяг, оскільки всі залежності гарантовано й автоматично встановлюються командою npm install на будь-якому сервері чи комп'ютері розробника.")

    add_placeholder(doc, "[ВСТАВИТИ ЗОБРАЖЕННЯ ЛІСТИНГУ 2.1: Конфігурація монорепозиторію та робочих просторів package.json]")

    add_paragraph(doc, "Серцевиною бекенду є конфігураційний файл cms/src/payload.config.ts. На основі програмно описаних у ньому схем і моделей адмін-панель Payload CMS генерується повністю автоматично, тому розробнику не потрібно витрачати час на її ручну верстку. Конфігурація реєструє колекції даних, наскрізні налаштування Globals, сучасний візуальний текстовий редактор Lexical, модуль оптимізації графіки Sharp та налаштовує автоматичний експорт TypeScript-типів у файл payload-types.ts. Для безперешкодної взаємодії фронтенду з сервером через різні порти чи домени увімкнено крос-доменні політики CORS (cors: '*').")

    add_paragraph(doc, "Питання вибору реляційної СУБД було безпосередньо пов'язане зі специфікою деплою на хмарну платформу Vercel. Оскільки Vercel функціонує за безсерверною (серверлес) архітектурою та має непостійну файлову систему (ephemeral storage), яка очищається при кожному оновленні чи перезапуску контейнерів, підняти реляційну базу локально (наприклад, у файлі SQLite) було неможливо — збережені дані втрачалися б. Тому як персистентне сховище було обрано зовнішню хмарну СУБД Neon PostgreSQL 16. Підключення налаштовано через адаптер @payloadcms/db-postgres із пулінгом з'єднань через PgBouncer та обов'язковим SSL-шифруванням (параметр sslmode=require), що гарантує надійність навіть при пікових навантаженнях.")

    add_placeholder(doc, "[ВСТАВИТИ ЗОБРАЖЕННЯ ЛІСТИНГУ 2.2: Конфігурація Payload CMS та підключення до Neon DB (payload.config.ts)]")

    add_paragraph(doc, "Для забезпечення повного контролю контенту через адмін-панель та підтримки додавання нових сторінок розроблено прозору структуру сутностей у Payload CMS:")

    add_bullet(doc, "Pages — центральна колекція керування сторінками сайту. Кожен документ містить назву (title), унікальний URL-ідентифікатор (slug), статус публікації (isPublished) та динамічне масив-поле layout (конструктор Page Builder), що дозволяє компонувати сторінку з довільного набору блоків;")
    add_bullet(doc, "Categories — рубрикатор тематичних напрямків всесвіту LEGO (Star Wars, Marvel Super Heroes, LEGO Ninjago, Classic Castle & Space, Деталі та аксесуари) з описами та бейджами;")
    add_bullet(doc, "Media — централізоване медіасховище для завантаження 3D-моделей (.obj, .mtl), текстур, банерів та промо-фотографій з автоматичною оптимізацією;")
    add_bullet(doc, "Users — захищені облікові записи адміністраторів та редакторів контенту із солінням та хешуванням паролів за сучасними стандартами;")
    add_bullet(doc, "Globals (HeaderSettings, FooterSettings) — наскрізні налаштування сайту (назва бренду, контактний телефон, посилання на офіційний Telegram-канал @grombricks, динамічний список пунктів навігації navItems та копірайт), які доступні глобально для всього веб-ресурсу.")

    add_placeholder(doc, "[ВСТАВИТИ ЗОБРАЖЕННЯ ЛІСТИНГУ 2.3: Опис схеми колекції Pages (Pages.ts)]")

    add_paragraph(doc, "Конструктор блоків сторінки (Page Builder у cms/src/blocks/). Кожен функціональний блок реалізовано у вигляді незалежної схеми Block, що реєструється у масиві PageBlocks. Завдяки цьому адміністратор може змінювати будь-який текст сторінки без втручання у вихідний код фронтенду:")

    add_bullet(doc, "HeroBlock — головний екран із текстовим бейджем «ПРЕМІАЛЬНІ LEGO КОЛЕКЦІЇ», заголовком, підзаголовком, кнопками швидкої навігації «3D ЕКСПОЗИЦІЯ» та якірним переходом «ПЕРЕВАГИ»;");
    add_bullet(doc, "FeaturesBlock — блок ключових переваг магазину GromBricks з динамічним масивом карток: для кожного пункту через випадаючий список select обирається іконка (shield, package, truck, sparkles, check), заголовок, бейдж та опис;");
    add_bullet(doc, "ModelViewerBlock — блок інтерактивної 3D-вітрини Three.js: дозволяє змінювати заголовок, підзаголовок та керувати списком 3D-моделей (Golden Ninja Lloyd, Kai Jungle, Cole Jungle), редагуючи назву фігурки, серію, тег рідкісності, опис та шляхи до файлів .obj/.mtl;");
    add_bullet(doc, "CategoryCardsBlock — презентаційні картки напрямків колекціонування (Мініфігурки, Деталі та Зброя, Star Wars & Ninjago, Рідкісні Раритети) із зображеннями та бейджами;");
    add_bullet(doc, "TextBlock — новий універсальний блок текстового контенту для базових інформаційних сторінок (Про нас, Доставка та оплата, Контакти тощо): містить бейдж, заголовок, підзаголовок, основний текст із підтримкою форматування списків та виділений акцентний блок-примітку (highlightBox);");
    add_bullet(doc, "BannerBlock — промо-банер для спеціальних оголошень з можливістю вибору колірного стилю через select;");
    add_bullet(doc, "CallToActionBlock — конверсійний блок із закликом долучитися до Telegram-спільноти GromBricks.")

    add_placeholder(doc, "[ВСТАВИТИ ЗОБРАЖЕННЯ ЛІСТИНГУ 2.4: Схема текстового блоку контенту (TextBlock.ts)]")

    add_paragraph(doc, "Для обробки зображень у бекенді задіяно бібліотеку Sharp. Під час завантаження медіафайлів сервер автоматично оптимізує їх та конвертує у сучасний формат WebP, що суттєво заощаджує мережевий трафік і прискорює завантаження на смартфонах.")

    add_paragraph(doc, "Стандартний REST API Payload CMS повертає глибоко вкладені службові об'єкти з надлишковими метаданими. Щоб передавати на фронтенд чистий, плоский та оптимізований JSON, було реалізовано спеціалізовані ендпоінти, зареєстровані у payload.config.ts (Таблиця 2.1).")

    add_table_title(doc, "Таблиця 2.1 - Кастомні API-ендпоінти сервера GromBricks")
    headers_t21 = ["HTTP-метод та URL", "Модуль обробника", "Призначення для лендінгу та сторінок", "Формат відповіді"]
    data_t21 = [
        ["GET /api/common", "commonEndpoint.ts", "Віддача наскрізних контактів, телефону, посилання на Telegram, пунктів меню navItems та копірайту з Globals.", "{ success: true, data: { header, footer } }"],
        ["GET /api/page?path=...", "pageEndpoint.ts", "Універсальне отримання структури будь-якої сторінки за slug: повертає масив блоків або 404 Not Found, якщо сторінка відсутня чи в чернетках.", "{ success: true, data: { title, slug, blocks[] } }"],
        ["POST /api/seed", "seedEndpoint.ts", "Первинне наповнення бази даних категоріями, головною сторінкою та базовими сторінками (about, delivery) у Pages.", "{ success: true, message }"]
    ]
    add_table(doc, headers_t21, data_t21, [1.7, 1.4, 2.5, 1.4])

    add_paragraph(doc, "Головним маршрутом динамічної взаємодії є ендпоінт GET /api/page?path=... Він нормалізує переданий шлях, шукає відповідний документ у базі Neon, перевіряє ознаку публікації isPublished, формує чистий плоский JSON із масивом блоків і повертає його фронтенду. Для неіснуючих або неопублікованих документів ендпоінт коректно повертає HTTP-статус 404.")

    add_placeholder(doc, "[ВСТАВИТИ ЗОБРАЖЕННЯ ЛІСТИНГУ 2.5: Реалізація універсального ендпоінту сторінок (pageEndpoint.ts)]")

    # -------------------------------------------------------------------------
    # 2.2. Клієнтська частина та стани UI
    # -------------------------------------------------------------------------
    add_heading_2(doc, "2.2. Клієнтська частина та стани UI: асинхронні запити до API, клієнтська маршрутизація, шаблонізація блоків, адаптивна верстка, обробка станів (завантаження / 404 / помилка / готовність).")

    add_paragraph(doc, "Клієнтську частину (директорія frontend/) спроєктовано як високопродуктивний веб-додаток. Свідома відмова від важких SPA-фреймворків на користь чистого Vanilla TypeScript, сучасного HTML5 та модульного CSS3 під керуванням збирача Vite дозволила усунути накладні витрати на віртуальний DOM та забезпечити практично миттєве завантаження першого екрана.")

    add_paragraph(doc, "Клієнтська маршрутизація (SPA-роутер у main.ts). Щоб користувач міг переглядати не лише головний лендінг, а й створені через адмінку базові інформаційні сторінки (наприклад, /about чи /delivery), у головному скрипті реалізовано легкий клієнтський роутер без перезавантаження сторінки:")

    add_bullet(doc, "функція getCurrentPageSlug() зчитує поточний маршрут із URL-шляху, параметрів запиту (?p=slug) або хешу;")
    add_bullet(doc, "функція navigateTo(url) оновлює історію переходів браузера за допомогою HTML5 History API (history.pushState) та ініціює плавне завантаження потрібної сторінки;")
    add_bullet(doc, "слухач події popstate забезпечує природну роботу кнопок «Назад» та «Вперед» у браузері;")
    add_bullet(doc, "на внутрішніх сторінках автоматично формуються навігаційні хлібні крихти (Breadcrumbs) із кнопкою швидкого повернення «← На головну».")

    add_paragraph(doc, "Життєвий цикл запиту та робота з API (api.ts). При відкритті сторінки функція init() надсилає паралельні запити: GET /api/common для наповнення шапки й підвалу та GET /api/page?path=... для отримання контенту поточної сторінки. Взаємодію реалізовано через fetch() із синтаксисом async/await. Усі мережеві виклики огорнуто в блоки try...catch: у разі відсутності зв'язку з бекендом функція підхоплює надійний локальний fallback-набір даних, запобігаючи аварійному збою вітрини.")

    add_placeholder(doc, "[ВСТАВИТИ ЗОБРАЖЕННЯ ЛІСТИНГУ 2.6: Мережевий модуль клієнта з обробкою помилок (api.ts)]")

    add_paragraph(doc, "Динамічна шапка сайту (Header) та підвал (Footer). На відміну від статичного коду, навігаційне меню шапки та мобільного меню генерується динамічно з масиву navItems, отриманого з Globals адмінки. Завдяки цьому додавання нового пункту меню (наприклад, посилання на створену сторінку «Доставка») здійснюється адміністратором за кілька секунд прямо в адмін-панелі.")

    add_paragraph(doc, "Шаблонізація блоків (фабрика шаблонів у renderBlocks.ts). Функція renderBlocks перебирає масив блоків, отриманих від API, і за полем blockType делегує формування HTML-розмітки відповідним модулям із каталогу blocks/:")

    add_bullet(doc, "renderHeroBlock (hero.ts) — формує перший екран із бейджем «ПРЕМІАЛЬНІ LEGO КОЛЕКЦІЇ», заголовком, підзаголовком, кнопкою «3D ЕКСПОЗИЦІЯ» (з іконкою куба) та якірною кнопкою «ПЕРЕВАГИ» (з іконкою щита якості ri-shield-star-line, яка плавно скролить екран до блоку переваг без спливаючих вікон);");
    add_bullet(doc, "renderFeaturesBlock (features.ts) — динамічно рендерить картки переваг магазину GromBricks з відповідними векторними іконками (shield, truck, package, sparkles, check);");
    add_bullet(doc, "renderModelViewerBlock (modelViewer.ts) — створює розмітку інтерактивної 3D-вітрини: контейнер Three.js, навігаційні таби перемикання моделей, бейдж, детальний опис та кнопки дій;");
    add_bullet(doc, "renderCategoryCardsBlock (categoryCards.ts) — будує презентаційні картки напрямків колекціонування;");
    add_bullet(doc, "renderTextBlock (text.ts) — формує розмітку інформаційних сторінок: акуратне розбиття тексту на абзаци, марковані списки та акцентний блок примітки highlightBox;");
    add_bullet(doc, "renderBannerBlock (banner.ts) — будує промо-банер акційних пропозицій;");
    add_bullet(doc, "renderCtaBlock (cta.ts) — створює конверсійний блок підписки на офіційний Telegram-канал.")

    add_placeholder(doc, "[ВСТАВИТИ ЗОБРАЖЕННЯ ЛІСТИНГУ 2.7: Шаблонізатор сторінки (renderBlocks.ts)]")

    add_paragraph(doc, "Інтерактивна 3D-вітрина на базі Three.js (modelViewer.ts). Головною родзинкою лендінгу є повноцінний 3D-переглядач мініфігурок. Після монтування розмітки в DOM функція initThreeViewer() налаштовує сцену Three.Scene, камеру PerspectiveCamera, антиаліасинг WebGLRenderer та контролер орбітального обертання OrbitControls. Освітлення сцени скомпоновано за принципом триточкового студійного світла (Key Light золотистого відтінку, Fill Light холодного синього та Rim Light контурного підсвічування), що ефектно підкреслює глянцеві та металізовані деталі мініфігурок LEGO. Моделі завантажуються динамічно за допомогою OBJLoader та MTLLoader із центруванням і масштабуванням на п'єдесталі. При кліку на будь-який таб (наприклад, перемиканні з Golden Ninja на Kai чи Cole) скрипт оновлює текстові описи з CMS та плавно перезавантажує відповідну 3D-геометрію.")

    add_placeholder(doc, "[ВСТАВИТИ ЗОБРАЖЕННЯ ЛІСТИНГУ 2.8: Ініціалізація сцени Three.js (modelViewer.ts)]")

    add_paragraph(doc, "Адаптивна верстка та стилістика (frontend/src/style.css). Оформлення базується на технологіях CSS Grid та Flexbox. Колірна палітра відтворює естетику преміальної космічної цеглини (#05070d, #070b16) із теплими золотими акцентами (#e6bf55, #f0cf7a). Розмір головного заголовка Hero було зменшено до гармонійного clamp(34px, 3.6vw, 46px), що забезпечує збалансовану типографіку. Завдяки медіа-запитам (@media на 1024px, 768px та 480px) сайт адаптується до будь-яких екранів: навігація на смартфонах ховається у плавне висувне меню (.mobile-nav-drawer), 3D-полотно змінює розмір через ResizeObserver, а текстові блоки сторінки вибудовуються у зручний потік читання.")

    add_paragraph(doc, "Обробка станів користувацького інтерфейсу (UI States). Для передбачуваності взаємодії інтерфейс має чіткий візуальний відгук на кожен стан життєвого циклу даних:")

    add_bullet(doc, "стан завантаження (Loading) — під час очікування відповіді API або завантаження важких 3D-файлів OBJ/MTL користувач бачить фірмовий анімований спінер у вигляді обертового кубика LEGO (.lego-spinner);");
    add_bullet(doc, "стан готовності (Ready/Success) — плавна поява сформованих блоків у DOM з підключенням інтерактивних обробників;")
    add_bullet(doc, "стан 404 Not Found — якщо користувач переходить за неіснуючим маршрутом або сторінка знята з публікації, виводиться стилізований екран 404 з повідомленням та кнопками повернення на головну або переходу в адмінку;")
    add_bullet(doc, "стан помилки (Error) — при критичному мережевому збої виводиться інформативне повідомлення замість аварійного падіння програми.")

    # -------------------------------------------------------------------------
    # 2.3. Типізація на TypeScript
    # -------------------------------------------------------------------------
    add_heading_2(doc, "2.3. Типізація на TypeScript: створення типів для структури сторінки та її блоків, використання Type Aliases, Union Types та Intersection Types, типізація функцій і API.")

    add_paragraph(doc, "Статична типізація мови TypeScript стала надійним інструментом забезпечення стабільності та передбачуваності роботи застосунку. На відміну від звичайного JavaScript, строга система типів унеможливлює помилки несумісності структур даних ще на етапі написання коду, гарантує інтелектуальне автодоповнення (IntelliSense) і робить взаємодію між сервером та вітриною регламентованою (файли cms/src/types/shop.ts та frontend/src/types.ts).")

    add_paragraph(doc, "Завдяки TypeScript у проєкті чітко типізовано всю структуру сторінок. Інтерфейс PageResponse описує метадані документа (title, slug, path) та масив blocks: PageBlock[], що унеможливлює передачу некоректного блоку на клієнтську вітрину.")

    add_paragraph(doc, "Використання псевдонімів типів (Type Aliases). Для усунення «магічних рядків» і підвищення читабельності коду оголошено семантичні псевдоніми:")

    add_bullet(doc, "базові типи: export type ID = string | number, export type Slug = string;");
    add_bullet(doc, "ідентифікатори 3D-моделей: export type ModelId = 'golden_ninja' | 'kai_jungle' | 'cole_jungle' | string;");
    add_bullet(doc, "тематичні серії LEGO: export type LegoSeries = 'Star Wars' | 'Marvel Super Heroes' | 'Ninjago' | 'Classic Space' | 'Castle' | 'Technic';");
    add_bullet(doc, "стилістичні параметри та режими: export type Currency = 'UAH' | 'EUR' | 'USD', export type BannerVariant = 'legoYellow' | 'darkBrick' | 'accentRed', export type ViewerControlMode = 'orbit' | 'autoRotate' | 'static'.")

    add_paragraph(doc, "Однією з найпотужніших можливостей типізації стало використання дискримінованих об'єднань типів (Discriminated Unions) зі спільним маркерним полем blockType:")

    add_bullet(doc, "BlockType — дискримінант типів контентних блоків лендінгу та сторінок: 'hero' | 'features' | 'modelViewer' | 'categoryCards' | 'banner' | 'cta' | 'text';");
    add_bullet(doc, "PageBlockData — дискриміноване об'єднання типів усіх блоків: HeroBlockData | FeaturesBlockData | ModelViewerBlockData | CategoryCardsBlockData | BannerBlockData | CtaBlockData | TextBlockData. Завдяки полю blockType компілятор TypeScript автоматично звужує тип (type narrowing) усередині switch-блоку, гарантуючи коректність переданих властивостей;");
    add_bullet(doc, "UIState<T> — типізація життєвого циклу UI: { status: 'idle' } | { status: 'loading' } | { status: 'success'; data: T } | { status: 'empty' } | { status: 'error'; error: string };");
    add_bullet(doc, "ApiResponse<T> — єдиний контракт клієнт-серверного обміну: { success: true; data: T } | { success: false; error: string }.")

    add_placeholder(doc, "[ВСТАВИТИ ЗОБРАЖЕННЯ ЛІСТИНГУ 2.9: Дискриміноване об'єднання блоків та типізація TextBlock (shop.ts)]")

    add_paragraph(doc, "Для гнучкого розширення базових описів сутностей без дублювання властивостей застосовано типи перетину (Intersection Types):")

    add_bullet(doc, "ModelWithSceneSettings = Model3DItem & { initialCameraDistance?: number; rotationSpeed?: number } — доповнює 3D-модель специфічними параметрами сцени;");
    add_bullet(doc, "FeaturedCategoryCard = CategoryCardItem & { featuredUntil?: string; priorityOrder: number } — розширює картку пріоритетом показу;");
    add_bullet(doc, "PageWithBlocks = PageLayout & { lastUpdatedFormatted?: string } — доповнює макет сторінки відформатованою датою оновлення.")

    add_paragraph(doc, "Модулі modelViewer.ts, renderBlocks.ts та сервіси API містять суворі сигнатури функцій (наприклад, фільтрація моделей filterModels(models: Model3DItem[], options: ModelFilterOptions): Model3DItem[] та перемикач активної моделі SwitchModelHandler = (modelId: string) => Promise<boolean>), що унеможливлює передачу некоректних аргументів.")

    add_paragraph(doc, "Система Payload CMS побудована на TypeScript і підтримує автоматичну синхронізацію типів: за допомогою команди payload generate:types на основі конфігурацій колекцій генерується файл payload-types.ts, що гарантує безшовну синхронізацію бекенду та фронтенду.")

    # -------------------------------------------------------------------------
    # 2.4. Деплой та контрольний приклад
    # -------------------------------------------------------------------------
    add_heading_2(doc, "2.4. Деплой та контрольний приклад: розгортання проєкту на Vercel та Neon, оформлення репозиторію (README.md), контрольний приклад взаємодії.")

    add_paragraph(doc, "Для публічної демонстрації роботи системи було налаштовано автоматичне хмарне розгортання на платформі Vercel у зв'язці з базою даних Neon PostgreSQL. Для розгортання монорепозиторію на хостингу Vercel було підключено єдиний GitHub-репозиторій та налаштовано два окремі проєкти:")

    add_bullet(doc, "grombricks-cms (Бекенд) — Root Directory: cms, середовище збірки Next.js; у налаштуваннях середовища Vercel задано змінні оточення (Environment Variables): DATABASE_URL (рядок підключення до Neon із параметром sslmode=require) та PAYLOAD_SECRET;");
    add_bullet(doc, "grombricks-front (Вітрина лендінгу) — Root Directory: frontend, команда збірки npm run build, каталог виходу dist; змінна оточення VITE_API_BASE вказує на публічну адресу розгорнутого бекенду CMS;");
    add_bullet(doc, "безпека конфіденційних даних — локальні файли .env не потрапляють у відкритий git-репозиторій, тому всі секретні ключі та адреси бази даних налаштовуються виключно через зашифровану панель Environment Variables у Vercel.")

    add_paragraph(doc, "У кореневій папці монорепозиторію підготовлено чіткий інструктивний файл README.md. Він містить опис системних вимог (Node.js 20.18+, Neon PostgreSQL), структуру каталогів cms/ та frontend/, команди встановлення залежностей npm install, скрипти локального запуску (npm run dev:cms, npm run dev:frontend) та інструкцію із розгортання.")

    add_placeholder(doc, "[ВСТАВИТИ ЗОБРАЖЕННЯ ЛІСТИНГУ 2.10: Документація репозиторію README.md]")

    add_paragraph(doc, "Працездатність розробленої Fullstack-системи перевірено від лиця адміністратора (редагування контенту й створення сторінок) та відвідувача сайту.")

    add_paragraph(doc, "Сценарій 1. Робота адміністратора в системі Payload CMS:")

    add_bullet(doc, "1. Адміністратор переходить за адресою /admin, вводить логін і пароль та потрапляє до згенерованої панелі керування. На головному екрані доступний список сутностей: Pages, Categories, Media, Users та Globals.");
    add_placeholder(doc, "[ВСТАВИТИ РИСУНОК 2.7: Панель керування Payload CMS: головний екран сутностей]");

    add_bullet(doc, "2. Редагування будь-якого тексту на головній сторінці: адміністратор відкриває розділ Pages, обирає документ «home» та може змінити будь-яке слово у блоках Hero, Features чи 3D Вітрини (наприклад, відредагувати заголовок, змінити опис чи замінити бейдж). Після натискання кнопки «Save» дані фіксуються в Neon PostgreSQL і миттєво відображаються на сайті без перезапуску коду фронтенду.");
    add_placeholder(doc, "[ВСТАВИТИ РИСУНОК 2.8: Редагування блоків головної сторінки в колекції Pages]");

    add_bullet(doc, "3. Створення нової базової сторінки через CMS: адміністратор натискає «Create New» у колекції Pages, задає назву «Доставка та оплата», слаг «delivery», обирає текстовий блок TextBlock, заповнює текст умов доставки та акцентну примітку й публікує сторінку.");
    add_placeholder(doc, "[ВСТАВИТИ РИСУНОК 2.9: Створення нової інформаційної сторінки через TextBlock]");

    add_bullet(doc, "4. Прив'язка до навігації: у розділі Globals (HeaderSettings) адміністратор додає новий пункт у список navItems із назвою «ДОСТАВКА» та адресою «/delivery». Пункт автоматично з'являється у шапці та мобільному меню вітрини.");
    add_placeholder(doc, "[ВСТАВИТИ РИСУНОК 2.10: Керування пунктами меню в Globals (HeaderSettings)]");

    add_paragraph(doc, "Сценарій 2. Сценарій взаємодії відвідувача з клієнтським веб-додатком:")

    add_bullet(doc, "1. Відвідувач відкриває головну сторінку сайту. Завдяки оптимізації Vite сторінка завантажується практично миттєво. На першому екрані користувач бачить лаконічний бейдж «ПРЕМІАЛЬНІ LEGO КОЛЕКЦІЇ» (без зайвих символів), виважений заголовок «ЕКСКЛЮЗИВНІ ФІГУРКИ ТА ДЕТАЛІ LEGO», кнопку швидкого переходу до 3D-вітрини та якірну кнопку «ПЕРЕВАГИ» з іконкою щита якості;");
    add_placeholder(doc, "[ВСТАВИТИ РИСУНОК 2.11: Головний екран Hero з оновленою типографікою та іконками кнопок]");

    add_bullet(doc, "2. При натисканні на кнопку «ПЕРЕВАГИ» сторінка плавно скролиться до розділу стандартів якості без будь-яких сторонніх модальних вікон;");
    add_bullet(doc, "3. В інтерактивній 3D-вітрині користувач обертає фігурку золотого ніндзя Lloyd на 360°, масштабує деталі коліщатком миші під студійним світлом та перемикає таби на Кая і Коула, миттєво спостерігаючи оновлення 3D-геометрії та текстового опису з CMS;");
    add_placeholder(doc, "[ВСТАВИТИ РИСУНОК 2.12: Інтерактивний огляд 3D-моделей мініфігурок Three.js]");

    add_bullet(doc, "4. Клікнувши в меню на створену сторінку «ДОСТАВКА», клієнтський SPA-роутер плавно завантажує вміст сторінки через /api/page?path=/delivery без повного перезавантаження сайту, показуючи хлібні крихти з кнопкою швидкого повернення «← На головну»;");
    add_placeholder(doc, "[ВСТАВИТИ РИСУНОК 2.13: Відображення створеної сторінки «Доставка та оплата» з хлібними крихтами]");

    add_bullet(doc, "5. Перевірка обробки помилок: при переході за неіснуючою адресою (наприклад, /non-existent) інтерфейс показує стилізований екран 404 Not Found з кнопкою повернення на головну.")
    add_placeholder(doc, "[ВСТАВИТИ РИСУНОК 2.14: Стилізований екран помилки 404 Not Found]");

    # =========================================================================
    # ВИСНОВОК ДО РОЗДІЛУ 2
    # =========================================================================
    add_heading_1(doc, "Висновок до розділу 2")

    add_paragraph(doc, "У другому розділі звіту спроєктовано та практично реалізовано презентаційний Fullstack-лендінг бренду GromBricks за архітектурою монорепозиторію. Використання безголової Payload CMS 3.x у поєднанні з хмарною СУБД Neon PostgreSQL дозволило створити зручний блоковий конструктор Page Builder для гнучкого редагування абсолютно всіх текстів лендінгу безпосередньо через адмін-панель без втручання у вихідний код фронтенду, а також впровадити можливість легкого додавання нових базових інформаційних сторінок за допомогою спеціалізованого блоку TextBlock.")

    add_paragraph(doc, "Клієнтська вітрина на чистому Vanilla TypeScript під збіркою Vite відзначається високою швидкістю завантаження, повною адаптивністю, вражаючою інтерактивною 3D-вітриною Three.js зі студійним світлом, власним легким SPA-роутером із підтримкою хлібних крихт та надійною обробкою життєвого циклу станів інтерфейсу (Loading, Ready, 404 Not Found, Error). Застосування комплексної системи типів TypeScript (Type Aliases, Discriminated Unions, Intersection Types) усунуло ризики помилок несумісності даних, а успішне хмарне розгортання проєкту на платформі Vercel підтвердило надійність, гнучкість та практичну цінність створеного рішення.")

def run_landing_update():
    base_file = "ЗВІТ_3.docx"
    output_file = "ЗВІТ_3_оновлений.docx"

    if not os.path.exists(base_file):
        print(f"Error: {base_file} not found!")
        return

    print(f"Loading {base_file}...")
    doc = docx.Document(base_file)

    body = doc._body._element
    start_idx = None
    end_idx = None

    for idx, child in enumerate(body):
        text = child.text if hasattr(child, 'text') else ''
        if not text and hasattr(child, 'xpath'):
            t_elems = child.xpath('.//w:t')
            text = ''.join([t.text for t in t_elems if t.text])

        # Знаходимо початок Розділу 2
        if "РОЗДІЛ 2." in text or "РОЗДІЛ 2" in text:
            if start_idx is None:
                start_idx = idx

        # Знаходимо початок Розділу 3
        if "РОЗДІЛ 3." in text or "РОЗДІЛ 3" in text:
            if end_idx is None and start_idx is not None:
                end_idx = idx

    print(f"Found Section 2 range: element {start_idx} to {end_idx}")

    if start_idx is None or end_idx is None:
        print("Could not find start or end of Section 2!")
        return

    temp_doc = docx.Document()
    generate_clean_landing_section_2(temp_doc)

    target_element = body[end_idx]

    # Вставляємо нові елементи перед початком Розділу 3
    for child in temp_doc._body._element:
        new_child = copy.deepcopy(child)
        target_element.addprevious(new_child)

    # Збираємо старі елементи Розділу 2 для видалення
    elements_to_delete = []
    current = body[start_idx]
    while current != target_element:
        elements_to_delete.append(current)
        current = current.getnext()

    num_new_elements = len(temp_doc._body._element)
    old_elements = elements_to_delete[:-num_new_elements]
    print(f"Removing {len(old_elements)} old Section 2 elements...")
    for elem in old_elements:
        body.remove(elem)

    print(f"Saving updated document to {output_file}...")
    doc.save(output_file)
    print(f"Successfully saved {output_file}!")

    # Також пробуємо зберегти в ЗВІТ_3.docx, якщо файл не заблокований у Word
    try:
        doc.save("ЗВІТ_3.docx")
        print("Successfully updated ЗВІТ_3.docx as well!")
    except Exception as e:
        print("Note: ЗВІТ_3.docx is currently open in Word. Changes are in ЗВІТ_3_оновлений.docx.")

if __name__ == "__main__":
    run_landing_update()
