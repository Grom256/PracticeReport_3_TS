# -*- coding: utf-8 -*-
"""
Full script to append Section 2 to GromBricks Practice Report docx.
"""
import os
import sys
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_heading_1(doc, text):
    p = doc.add_paragraph(style='Heading 1')
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.first_line_indent = Inches(0)
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(14)
    run.font.bold = True
    return p

def add_heading_2(doc, text):
    p = doc.add_paragraph(style='Heading 2')
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.first_line_indent = Inches(0.492)
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(14)
    run.font.bold = True
    return p

def add_subheading(doc, text):
    p = doc.add_paragraph(style='Normal')
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.first_line_indent = Inches(0.492)
    p.paragraph_format.line_spacing = 1.5
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(14)
    run.font.bold = True
    return p

def add_paragraph(doc, text):
    p = doc.add_paragraph(style='Normal')
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.first_line_indent = Inches(0.492)
    p.paragraph_format.line_spacing = 1.5
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(14)
    return p

def add_placeholder(doc, text):
    p = doc.add_paragraph(style='Normal')
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.first_line_indent = Inches(0)
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.italic = True
    return p

def add_caption(doc, text):
    p = doc.add_paragraph(style='No Spacing')
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.first_line_indent = Inches(0)
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    run.font.italic = True
    return p

def add_table_title(doc, text):
    p = doc.add_paragraph(style='Normal')
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.first_line_indent = Inches(0)
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    return p

def add_table(doc, headers, data, col_widths=None):
    table = doc.add_table(rows=len(data) + 1, cols=len(headers), style='Table Grid')
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    hdr_cells = table.rows[0].cells
    for i, header_text in enumerate(headers):
        hdr_cells[i].text = header_text
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=150, right=150)
        p = hdr_cells[i].paragraphs[0]
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Inches(0)
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        for r in p.runs:
            r.font.name = 'Times New Roman'
            r.font.size = Pt(11)
            r.font.bold = True

    for row_idx, row_data in enumerate(data):
        row_cells = table.rows[row_idx + 1].cells
        for col_idx, cell_value in enumerate(row_data):
            row_cells[col_idx].text = cell_value
            set_cell_margins(row_cells[col_idx], top=100, bottom=100, left=150, right=150)
            p = row_cells[col_idx].paragraphs[0]
            p.paragraph_format.first_line_indent = Inches(0)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            if col_idx == 0:
                p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            else:
                p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            for r in p.runs:
                r.font.name = 'Times New Roman'
                r.font.size = Pt(11)

    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Inches(w)

    doc.add_paragraph(style='Normal').paragraph_format.space_after = Pt(4)
    return table

def build_section_2(doc):
    # =========================================================================
    # РОЗДІЛ 2
    # =========================================================================
    add_heading_1(doc, "РОЗДІЛ 2. РОЗРОБКА FULLSTACK-ЗАСТОСУНКУ НА PAYLOAD CMS ТА TYPESCRIPT (ГУСАРОВ, ЮРЧИШИН)")

    # -------------------------------------------------------------------------
    # 2.1. Архітектура та Payload CMS
    # -------------------------------------------------------------------------
    add_heading_2(doc, "2.1. Архітектура та Payload CMS: монорепозиторій, конфігурація Payload (колекції, типи блоків, медіа), хмарна БД та кастомні API-ендпоінти.")

    add_paragraph(doc, "Другий етап навчальної практики присвячений проєктуванню та розробці сучасного Fullstack-застосунку інтернет-магазину GromBricks із застосуванням інноваційного стеку технологій: безголової системи керування контентом Payload CMS 3.x, платформи Next.js, мови програмування TypeScript, інструменту збірки Vite, хмарної реляційної бази даних Neon PostgreSQL та хостингу Vercel. На відміну від класичного монолітного підходу, використаного в першому розділі на базі Django MVT, тут реалізовано decoupled-архітектуру з чітким розділенням відповідальності між бекендом (Headless CMS) та клієнтською вітриною магазину.")

    add_paragraph(doc, "Для ефективної організації кодової бази проєкту застосовано концепцію монорепозиторію (monorepo). Монорепозиторій об'єднує дві автономні підсистеми — бекенд (директорія cms/) та фронтенд (директорія frontend/) — у межах єдиного Git-репозиторію з централізованим керуванням залежностями за допомогою механізму npm workspaces. Кореневий файл package.json містить декларацію робочих просторів [\"cms\", \"frontend\"] та набір уніфікованих скриптів життєвого циклу: dev:cms для запуску середовища розробки Payload CMS на порту 3000, dev:frontend для запуску локального сервера Vite на порту 5173, build:cms та build:frontend для автономної компіляції компонентів, а також підсумковий скрипт build для конвеєрного складання всього проєкту.")

    add_paragraph(doc, "Така організація усуває необхідність синхронізації кількох окремих репозиторіїв, забезпечує спільне повторне використання типів даних TypeScript та спрощує налаштування неперервної інтеграції й деплою (CI/CD). Архітектура монорепозиторію має таку ієрархічну структуру:")

    add_paragraph(doc, "PractiseReport3_ts/\n"
                       "├── cms/                  # Бекенд: Payload CMS 3.x + Next.js App Router + TypeScript\n"
                       "│   ├── src/\n"
                       "│   │   ├── collections/  # Схеми колекцій: Products, Categories, Orders, Pages, Media, Users\n"
                       "│   │   ├── globals/      # Глобальні сутності: HeaderSettings, FooterSettings\n"
                       "│   │   ├── blocks/       # Декларації блоків сторінки: Hero, Features, ProductGrid, Banner, CTA\n"
                       "│   │   ├── endpoints/    # Кастомні API-ендпоінти: common, page, search, cart, checkout, seed\n"
                       "│   │   ├── types/        # Суворі типи предметної області магазину (shop.ts)\n"
                       "│   │   ├── lib/          # Сервісні бізнес-функції валідації та розрахунків (shopUtils.ts)\n"
                       "│   │   └── payload.config.ts # Головна конфігурація CMS, бази даних, CORS та плагінів\n"
                       "│   └── package.json\n"
                       "│\n"
                       "└── frontend/             # Клієнтська вітрина: Vite + Vanilla TypeScript + HTML5 / CSS3\n"
                       "    ├── index.html        # Головна сторінка (інтерактивний лендінг та 3D-експозиція)\n"
                       "    ├── catalog.html      # Сторінка каталогу товарів із живим пошуком та фільтрами\n"
                       "    ├── cart.html         # Сторінка кошика з динамічним розрахунком суми\n"
                       "    ├── checkout.html     # Сторінка оформлення замовлення з клієнтською валідацією\n"
                       "    ├── src/\n"
                       "    │   ├── api.ts        # Сервісний шар взаємодії з бекендом через fetch API\n"
                       "    │   ├── cart.ts       # Модуль керування кошиком та збереження у localStorage\n"
                       "    │   ├── blocks/       # Шаблонізатори блоків сторінки (renderBlocks.ts, hero.ts тощо)\n"
                       "    │   └── style.css     # Преміальна темна дизайн-система із золотими акцентами\n"
                       "    ├── vite.config.ts    # Конфігурація збірки та локального проксіювання API\n"
                       "    └── package.json")

    add_placeholder(doc, "[ВСТАВИТИ ЗОБРАЖЕННЯ ЛІСТИНГУ 2.1: Конфігурація монорепозиторію та робочих просторів package.json]")

    add_paragraph(doc, "Важливим архітектурним рішенням стала інтеграція хмарної реляційної СУБД Neon PostgreSQL замість традиційних локальних баз даних (SQLite або MySQL). Необхідність використання зовнішньої хмарної бази даних зумовлена специфікою цільової платформи розгортання Vercel. Платформа Vercel функціонує на базі безсерверних функцій (Serverless Functions) з ефемерною (тимчасовою) файловою системою. Будь-які локальні файли даних (наприклад, файли sqlite.db або завантажені користувачами зображення) безповоротно втрачаються під час автоматичного перезапуску або нового деплою безсерверного контейнера. Neon надає повністю керовану, масштабовану хмарну PostgreSQL із підтримкою пулінгу з'єднань через PgBouncer. Підключення до Neon налаштовано у файлі payload.config.ts за допомогою адаптера @payloadcms/db-postgres із параметром захищеного з'єднання ssl: { rejectUnauthorized: false }.")

    add_paragraph(doc, "Головний конфігураційний файл cms/src/payload.config.ts визначає всі ключові параметри функціонування системи: налаштування точки входу в панель адміністрування (/admin), користувача автентифікації Users.slug, підключення сучасного текстового редактора Lexical (@payloadcms/richtext-lexical), бібліотеки обробки зображень Sharp та вихідного шляху генерації типів TypeScript payload-types.ts. Для забезпечення взаємодії між клієнтською вітриною та сервером у конфігурації активовано директиву cors: '*' та очищено обмеження csrf. Це уможливлює безперешкодне виконання асинхронних cross-origin запитів від фронтенду до CMS як під час локальної розробки, так і у продуктовому середовищі на платформі Vercel.")

    add_placeholder(doc, "[ВСТАВИТИ ЗОБРАЖЕННЯ ЛІСТИНГУ 2.2: Конфігурація Payload CMS та підключення хмарної бази даних (payload.config.ts)]")

    add_paragraph(doc, "Структура даних у Payload CMS будується на основі декларативних колекцій (Collections) та глобальних об'єктів (Globals). Кожна колекція відповідає за окрему сутність бізнес-домену магазину LEGO GromBricks:")

    add_paragraph(doc, "1. Колекція Products (Товари) — центральна сутність каталогу. Вона містить текстове поле назви (name), унікальний URL-ідентифікатор (slug), обов'язковий унікальний артикул (sku), селектор типу продукту productType зі значеннями minifigure (колекційна фігурка) та part (окрема деталь), селектор ліцензійної серії series (Star Wars, Marvel Super Heroes, Ninjago, Classic Space, Castle тощо), числове поле вартості у гривнях (price), оцінку стану condition (новий або вживаний), залишок на складі у штуках (stock), булевий прапорець рідкісного ексклюзивного товару (isRare), реляційний зв'язок relationship із колекцією categories, поле завантаженого медіа images та альтернативне текстове посилання imageUrl, детальний опис (description), колір деталі (color), заводський Lego Element ID (partNumber) та статус активності (isActive). Права доступу на читання налаштовано як публічні (read: () => true).")

    add_paragraph(doc, "2. Колекція Categories (Категорії) — забезпечує структуризацію асортименту за всесвітами та типами конструктора. Включає назву категорії (title), унікальний slug, опис (description), текстовий бейдж (badge) для маркування («Хіт», «Новинка», «Топ») та назву векторної іконки (iconName).")

    add_paragraph(doc, "3. Колекція Orders (Замовлення) — фіксує результати покупок клієнтів. Кожне замовлення характеризується автоматично згенерованим номером (orderNumber виду GB-2026-XXXX), контактними відомостями клієнта (customerName, customerEmail, customerPhone), обраною службою доставки (deliveryService: Нова Пошта, Укрпошта або Meest), населеним пунктом (deliveryCity) та номером відділення/поштомату (deliveryBranch). Блок фінансів містить спосіб оплати paymentMethod (переказ на IBAN/картку або накладений платіж при отриманні), поточний стан обробки status (новий pending, в обробці processing, відправлений shipped, виконаний delivered або скасований cancelled) та фінальну суму totalAmount. Склад замовлення зберігається у масиві items, де кожен елемент містить посилання на товар product, назву товару, артикул sku, зафіксовану на момент покупки ціну одиниці unitPrice та кількість quantity. Для цієї колекції відкрито публічний доступ на створення записів (create: () => true), що дає змогу оформлювати замовлення безпосередньо з вітрини магазину.")

    add_paragraph(doc, "4. Колекція Pages (Сторінки) — реалізує модульний візуальний конструктор сторінок (Page Builder). Документ сторінки має заголовок (title), slug (наприклад, home для головної сторінки), прапорець публікації (isPublished) та поліморфне масивне поле layout типу blocks. Редактор може динамічно додавати, видаляти та змінювати порядок контентних блоків прямо в адмін-панелі.")

    add_paragraph(doc, "5. Колекція Media (Медіафайли) — керує завантаженням графічних матеріалів. Налаштування upload задають каталог збереження public/media, підтримку mime-типів зображень та автоматичну генерацію трьох оптимізованих розмірів бібліотекою Sharp: thumbnail (300×300 px), card (600×600 px) та hero (1440×600 px).")

    add_paragraph(doc, "6. Колекція Users (Користувачі) — забезпечує автентифікацію в системі керування контентом. Підтримує рольову модель role із рівнями доступу: admin (повні права), editor (редагування каталогу) та customer (клієнт).")

    add_table_title(doc, "Таблиця 2.1 - Основні колекції та глобальні налаштування Payload CMS")
    add_table(doc,
              ["Сутність / Slug", "Тип у Payload", "Призначення та основні поля", "Правила доступу"],
              [
                  ["products", "Collection", "Каталог деталей та мініфігурок LEGO: name, slug, sku, productType, series, price, condition, stock, isRare, category, imageUrl, isActive.", "read: public, write: admin/editor"],
                  ["categories", "Collection", "Тематичні класифікатори асортименту: title, slug, description, badge, iconName.", "read: public, write: admin"],
                  ["orders", "Collection", "Історія та деталі замовлень: orderNumber, customerName, customerEmail, customerPhone, deliveryService, status, totalAmount, items[].", "read: public, create: public, update: admin"],
                  ["pages", "Collection", "Сторінки з динамічним блоковим конструктором: title, slug, isPublished, layout (масив blocks).", "read: public, write: admin/editor"],
                  ["media", "Collection", "Завантаження та нарізка зображень Sharp: alt, caption, imageSizes (thumbnail, card, hero).", "read: public, upload: admin/editor"],
                  ["users", "Collection", "Системні користувачі з підтримкою JWT-auth: name, email, role (admin / editor / customer), phone.", "auth: true, read: public, manage: admin"],
                  ["header-settings", "Global", "Наскрізні налаштування шапки: brandName, phone, telegramUrl, navItems[].", "read: public, write: admin"],
                  ["footer-settings", "Global", "Наскрізні налаштування підвалу: description, email, copyright.", "read: public, write: admin"]
              ],
              col_widths=[1.4, 1.1, 3.2, 1.3])

    add_placeholder(doc, "[ВСТАВИТИ ЗОБРАЖЕННЯ ЛІСТИНГУ 2.3: Схема колекції Products та визначення полів (Products.ts)]")

    add_paragraph(doc, "Концепція побудови сторінок на основі блоків (Block-based CMS) реалізована за допомогою масиву PageBlocks у директорії cms/src/blocks/. Кожен блок є ізольованим модулем із власною схемою полів:")
    add_paragraph(doc, "• HeroBlock (slug: 'hero') — вступний банер із полями heading, subheading, badge («100% Оригінал LEGO®»), посиланнями та текстом головної й другорядної кнопок, а також фоновим зображенням imageUrl;\n"
                       "• FeaturesBlock (slug: 'features') — блок переваг із динамічним масивом items, де для кожного пункту обирається векторна іконка (щит, вантажівка, іскри, пакунок, галочка), назва та детальний опис;\n"
                       "• ProductGridBlock (slug: 'productGrid') — товарна сітка з можливістю конфігурації заголовка, підзаголовка, ліміту вибірки (від 2 до 24) та правила фільтрації filterType (all, minifigures, parts, rare);\n"
                       "• BannerBlock (slug: 'banner') — акційний блок зі знижками, текстом заклику до дії та колірною схемою variant (legoYellow, darkBrick, accentRed);\n"
                       "• CallToActionBlock (slug: 'cta') — блок залучення клієнтів із посиланням на Telegram-спільноту магазину.")

    add_paragraph(doc, "Глобальні сутності HeaderSettings та FooterSettings (каталог cms/src/globals/) використовуються для збереження контенту, який повторюється на всіх сторінках сайту. Вони містять офіційний контактний телефон, посилання на службу підтримки у месенджері Telegram, перелік навігаційних посилань, опис місії магазину та копірайт.")

    add_paragraph(doc, "Стандартний REST API Payload CMS повертає глибоко вкладені JSON-відповіді з надлишковою службовою інформацією, метаданими пагінації та нерозгорнутими зв'язками. Це створює додаткове навантаження на мережевий трафік та суттєво ускладнює парсинг на клієнті без важких бібліотек. Для оптимізації взаємодії з клієнтською частиною було реалізовано патерн Backend for Frontend (BFF) у вигляді кастомних API-ендпоінтів у директорії cms/src/endpoints/ та підключено їх до Next.js App Router (cms/src/app/api/):")

    add_paragraph(doc, "1. GET /api/common — зчитує глобальні сутності HeaderSettings та FooterSettings за допомогою req.payload.findGlobal() і формує плоский об'єкт із навігацією та контактами. У разі відсутності записів у базі активується вбудований аварійний fallback-об'єкт.")

    add_paragraph(doc, "2. GET /api/page?path=/ — обробляє запит на отримання структури сторінки. Знаходить документ у колекції pages за полем slug. Якщо сторінка містить блок типу productGrid, ендпоінт автоматично робить вибірку з 12 актуальних активних товарів із колекції products і вбудовує їх безпосередньо у відповідний блок сторінки. У результаті фронтенд отримує повністю готовий до відображення плоский масив блоків за один HTTP-запит.")

    add_paragraph(doc, "3. GET /api/products/search — реалізує високопродуктивний повнотекстовий пошук та фільтрацію каталогу. Приймає URL-параметри: q (пошуковий запит), type (minifigure/part), series (серія LEGO), condition (стан), onlyRare (лише рідкісні) та sortBy (сортування за ціною чи назвою).")

    add_paragraph(doc, "4. POST /api/cart/validate — здійснює валідацію кошика перед покупкою. Порівнює масив productId та quantity з актуальними залишками на складі в базі даних Neon, запобігаючи замовленню відсутніх позицій.")

    add_paragraph(doc, "5. POST /api/orders/checkout — критично важливий транзакційний ендпоінт. Перевіряє повноту контактних даних клієнта та адресних реквізитів доставки, валідує актуальні залишки всіх замовлених товарів, генерує унікальний номер замовлення вида GB-2026-XXXX, створює запис у колекції orders та атомарно зменшує поле stock у відповідних документах колекції products.")

    add_paragraph(doc, "6. POST /api/seed — сервісний ендпоінт автоматизованого початкового заповнення бази даних. Створює 5 базових категорій (Star Wars, Marvel, Ninjago, Classic, Parts), 8 культових товарів (Дарт Вейдер з принтом рук, Клон-капітан Рекс, Людина-Павук, Залізна Людина, світлові мечі тощо) та формує головну сторінку з повним набором динамічних блоків.")

    add_table_title(doc, "Таблиця 2.2 - Кастомні API-ендпоінти сервера GromBricks")
    add_table(doc,
              ["HTTP-метод та URL", "Модуль обробника", "Вхідні параметри", "Формат та опис відповіді"],
              [
                  ["GET /api/common", "commonEndpoint.ts", "Відсутні", "{ success: true, data: { header, footer } } — глобальні налаштування сайту."],
                  ["GET /api/page", "pageEndpoint.ts", "path (шлях до сторінки, наприклад / або /about)", "{ success: true, data: { title, slug, blocks[] } } — плоскі блоки з товарами."],
                  ["GET /api/products/search", "searchEndpoint.ts", "q, type, series, condition, onlyRare, minPrice, maxPrice, sortBy", "{ success: true, data: Product[] } — відфільтрований перелік позицій каталогу."],
                  ["POST /api/cart/validate", "cartEndpoints.ts", "JSON: { items: [{ productId, quantity }] }", "{ success: true, data: CartValidationResult } — перевірка наявності та актуальності цін."],
                  ["POST /api/orders/checkout", "orderEndpoints.ts", "JSON: { customer, delivery, paymentMethod, items, notes }", "{ success: true, data: { orderId, orderNumber, totalAmount } } — створення замовлення та списання залишку."],
                  ["POST /api/seed", "seedEndpoint.ts", "Відсутні (або ключ безпеки)", "{ success: true, message, categoriesCount, productsCount } — первинне наповнення БД."]
              ],
              col_widths=[1.7, 1.4, 1.8, 2.1])

    add_placeholder(doc, "[ВСТАВИТИ ЗОБРАЖЕННЯ ЛІСТИНГУ 2.4: Транзакційне оформлення замовлення у файлі orderEndpoints.ts]")

    # -------------------------------------------------------------------------
    # 2.2. Клієнтська частина та стани UI
    # -------------------------------------------------------------------------
    add_heading_2(doc, "2.2. Клієнтська частина та стани UI: асинхронні запити до API, шаблонізація блоків, адаптивна верстка, обробка станів (завантаження / помилка / готовність).")

    add_paragraph(doc, "Клієнтська вітрина інтернет-магазину GromBricks розроблена у директорії frontend/ відповідно до суворих методичних вимог: без використання важких реактивних компонентних фреймворків (таких як React, Angular чи Vue). Вітрина побудована на базі чистого Vanilla TypeScript, сучасного стандарту семантичної розмітки HTML5, модульного CSS3 та надшвидкого збирача проєктів Vite. Такий підхід забезпечує бездоганну швидкість первинного рендерингу (First Contentful Paint), нульовий оверхед стороннього віртуального DOM, мінімальний підсумковий розмір збірки та повний контроль над життєвим циклом DOM-елементів.")

    add_paragraph(doc, "Застосунок організовано за багатосторінковою моделлю (Multi-Page Application, MPA), що містить 4 повнофункціональні сторінки:")
    add_paragraph(doc, "1. index.html — презентаційний інтерактивний лендінг магазину. Містить іміджевий Hero-екран, блок переваг для колекціонерів AFOL (Adult Fans of LEGO), інтерактивну 3D-вітрину Three.js для огляду мініфігурок, картки ключових напрямків асортименту та блок заклику до дії (CTA).\n"
                       "2. catalog.html — функціональна вітрина каталогу. Забезпечує живий повнотекстовий пошук, фільтрацію за табами типів (Усі / Мініфігурки / Деталі), чекбокс відбору рідкісних ексклюзивів, сортування та миттєве додавання товарів до кошика.\n"
                       "3. cart.html — повноцінна сторінка кошика покупця. Відображає таблицю обраних позицій, мініатюри, ціни, елементи оперативного збільшення/зменшення кількості, кнопковий механізм видалення та зведений розрахунок вартості замовлення.\n"
                       "4. checkout.html — сторінка фіналізації замовлення. Включає перелік замовлених товарів, форми введення контактних даних покупця, вибору відділення перевізника, способу оплати та блок відображення фінального статусу.")

    add_paragraph(doc, "Організація мережевої взаємодії з бекендом винесена в сервісний модуль frontend/src/api.ts. Запити виконуються через нативний інтерфейс fetch із використанням синтаксису async/await. Базова адреса сервера формується динамічно за допомогою змінної оточення import.meta.env.VITE_API_BASE. У середовищі розробки Vite автоматично перенаправляє запити з префіксом /api на локальний бекенд Payload CMS (http://localhost:3000), а в продуктовому деплої звертається за публічною хмарною URL-адресою.")

    add_paragraph(doc, "Для забезпечення відмовостійкості всі асинхронні виклики (fetchCommonData, fetchPageData, fetchProducts, submitCheckout) загорнуті в блоки try...catch. У разі виникнення помилок мережі або тимчасової недоступності сервера бекенду система не аварійно припиняє роботу, а переходить у режим безперервної доступності (graceful degradation), повертаючи заздалегідь структуровані резервні дані (fallback).")

    add_placeholder(doc, "[ВСТАВИТИ ЗОБРАЖЕННЯ ЛІСТИНГУ 2.5: Асинхронні запити до API та механізм fallback-даних (api.ts)]")

    add_paragraph(doc, "Шаблонізація контенту побудована на патерні «Фабрика шаблонів» (Template Factory) у модулі frontend/src/blocks/renderBlocks.ts. Функція renderBlocks приймає масив блоків, отриманий від ендпоінту /api/page, виконує його ітерацію за допомогою методу map та викликає спеціалізований генератор розмітки відповідно до значення властивості blockType:")

    add_paragraph(doc, "• renderHeroBlock(block) — генерує верстку головного банера з динамічним підставленням заголовків, бейджа «PREMIUM LEGO COLLECTIBLES», кнопок дій та фонового зображення;\n"
                       "• renderFeaturesBlock(block) — формує адаптивну сітку карток переваг магазину з SVG-іконками з бібліотеки Remix Icon (щит, вантажівка, іскри, пакунок);\n"
                       "• renderCategoryCardsBlock() — рендерить візуальні картки тематичних колекцій (Star Wars, Marvel, Ninjago, Classic) із прямими переходами до каталогу;\n"
                       "• renderProductGridBlock(block) та renderProductCard(product) — формують товарну сітку з преміальними картками, що містять фотографію, статус наявності, артикул SKU, вартість та кнопку «В кошик» з data-атрибутами;\n"
                       "• renderCtaBlock(block) — верстає блок підписки на офіційну спільноту в Telegram;\n"
                       "• renderModelViewerBlock() — ініціалізує контейнер для інтерактивної 3D-сцени.")

    add_paragraph(doc, "Особливою родзинкою клієнтської вітрини є інтеграція бібліотеки Three.js у модулі modelViewer.ts. Модуль ініціалізує тривимірну сцену WebGL з динамічним спрямованим і розсіяним освітленням, орбітальним керуванням камерою (OrbitControls), обробкою жестів миші та сенсорного екрана, а також підтримкою завантаження й рендерингу оптимізованих 3D-моделей мініфігурок LEGO (Kai, Cole, Golden Ninja) у форматах OBJ/MTL та Collada DAE з плавним перемиканням між персонажами.")

    add_placeholder(doc, "[ВСТАВИТИ ЗОБРАЖЕННЯ ЛІСТИНГУ 2.6: Реалізація шаблонізатора сторінки в renderBlocks.ts]")

    add_paragraph(doc, "Адаптивна верстка та візуальна стилізація реалізовані у таблиці стилів frontend/src/style.css. Дизайн-система спирається на фірмову палітру темних відтінків космічної цегли (#05070d, #070b16, #0e1526) у поєднанні з контрастними преміальними золотими акцентами LEGO (#e6bf55, #f0cf7a, #ffd700). Завдяки використанню сучасної модульної сітки CSS Grid із правилом repeat(auto-fill, minmax(280px, 1fr)) елементи каталогу автоматично адаптуються під будь-яку ширину екрана без порушення візуальної ієрархії.")

    add_paragraph(doc, "Для підтримки мобільних пристроїв та планшетів визначено систему медіа-запитів (@media (max-width: 1024px), @media (max-width: 768px), @media (max-width: 480px)):")
    add_paragraph(doc, "• на смартфонах верхнє горизонтальне меню приховується та замінюється ергономічною кнопкою-бургером, яка активує анімовану бічну панель навігації (.mobile-nav-drawer);\n"
                       "• товарна сітка автоматично перебудовується з 4 колонок у 2 на планшетах та в 1 повнорозмірну колонку на мобільних пристроях;\n"
                       "• широка таблиця кошика трансформується у вертикальну послідовність карток товарів, зберігаючи зручний доступ до кнопок зміни кількості;\n"
                       "• форма оформлення замовлення перемикається з двоколонкової схеми в єдиний послідовний потік введення даних.")

    add_paragraph(doc, "Життєвий цикл користувацького інтерфейсу на кожній сторінці передбачає обов'язкову обробку чотирьох канонічних станів UI (UI States), що виключає появу неінформативних «білих екранів»:")
    add_paragraph(doc, "1. Стан очікування / завантаження (Loading State) — активується в момент надсилання асинхронного запиту до API. Контейнер контенту заповнюється фірмовим CSS-індикатором .lego-spinner, що візуалізує обертання деталі конструктора LEGO, та повідомленням («Завантаження GromBricks...» або «Пошук товарів...»).\n"
                       "2. Стан готовності / успіху (Success / Ready State) — настає після успішного одержання відповіді від сервера. DOM-контейнер миттєво наповнюється згенерованою структурою блоків або товарною сіткою, після чого підключаються слухачі подій кліків.\n"
                       "3. Стан порожніх даних (Empty State) — відображається у разі відсутності записів у відповіді (порожній кошик або відсутність товарів за введеним пошуковим фільтром). Користувачеві демонструється блок .empty-container з іконкою, роз'яснювальним текстом («За вашим запитом не знайдено жодної позиції LEGO») та кнопкою швидкого переходу до повного каталогу.\n"
                       "4. Стан помилки (Error State) — перехоплює мережеві збої або серверні коди помилок 4xx/5xx. Замість аварійної зупинки застосунку виводиться адаптивний банер .error-container або .checkout-error-banner із описом проблеми та можливістю повторити запит, а кнопка відправки форми блокується для запобігання дублюванню замовлень.")

    add_table_title(doc, "Таблиця 2.3 - Стани користувацького інтерфейсу вітрини магазину та їх обробка")
    add_table(doc,
              ["Стан інтерфейсу", "Умова виникнення", "Візуальний компонент та CSS-клас", "Поведінка системи та зворотний зв'язок"],
              [
                  ["Loading (Завантаження)", "Початок виконання fetch-запиту до ендпоінтів /api/page або /api/products/search.", ".loading-container, .lego-spinner, .loading-text", "Відображення фірмової анімації обертання кубика LEGO, блокування кліків."],
                  ["Success (Готовність)", "Отримання коду 200/201 та валідного JSON-масиву даних.", ".catalog-grid, .cart-page-grid, .confirmation-card", "Рендеринг карток товарів або блоків сторінки, ініціалізація слухачів подій."],
                  ["Empty (Порожньо)", "Пошуковий фільтр повернув 0 товарів або масив кошика містить 0 позицій.", ".empty-container, .empty-title, .empty-message", "Виведення дружнього тексту з пропозицією скинути фільтри або перейти в каталог."],
                  ["Error (Помилка)", "Мережевий збій, недоступність бекенду або помилка валідації 400 Bad Request.", ".error-container, .checkout-error-banner", "Інформативне повідомлення з текстом помилки, розблокування форми для виправлення."]
              ],
              col_widths=[1.5, 1.8, 1.8, 1.9])

    # -------------------------------------------------------------------------
    # 2.3. Типізація на TypeScript
    # -------------------------------------------------------------------------
    add_heading_2(doc, "2.3. Типізація на TypeScript: створення типів для структури сторінки та її блоків, використання Type Aliases, Union Types та Intersection Types для опису різних типів даних, типізація функцій і об'єктів, типізація даних, отриманих через API.")

    add_paragraph(doc, "Використання строгої статичної типізації мови TypeScript є фундаментальною основою надійності, передбачуваності та безпеки розроблюваного Fullstack-застосунку GromBricks. Типізація охоплює всі рівні програмної архітектури: схеми колекцій бекенду, опис динамічних сторінок, кошик покупця, алгоритми валідації та структури клієнтських запитів до API. Централізовані визначення типів винесено у файли cms/src/types/shop.ts та frontend/src/types.ts.")

    add_paragraph(doc, "Для створення виразної предметної моделі та позбавлення від небезпечних «магічних рядків» і довільних числових значень використано псевдоніми типів (Type Aliases):")
    add_paragraph(doc, "• export type ID = string | number — узагальнений ідентифікатор сутності, що забезпечує сумісність між різними діалектами баз даних;\n"
                       "• export type Slug = string — псевдонім для валідних URL-ідентифікаторів;\n"
                       "• export type ProductType = 'minifigure' | 'part' — обмеження можливих типів товарів виключно колекційними фігурками або окремими деталями;\n"
                       "• export type ProductCondition = 'new' | 'used' — допустимі стани збереження товару;\n"
                       "• export type OrderStatus = 'pending' | 'processing' | 'shipped' | 'delivered' | 'cancelled' — строго зафіксований перелік статусів життєвого циклу замовлення;\n"
                       "• export type DeliveryService = 'nova_poshta' | 'ukrposhta' | 'meest' — доступні логістичні оператори України;\n"
                       "• export type PaymentMethod = 'card_transfer' | 'cash_on_delivery' — дозволені методи проведення оплати;\n"
                       "• export type LegoSeries = 'Star Wars' | 'Marvel Super Heroes' | 'Ninjago' | 'Castle' | ... — строго типізований список офіційних серій LEGO;\n"
                       "• export type Currency = 'UAH' | 'EUR' | 'USD' — валюта обліку грошових сум;\n"
                       "• export type SortOption = 'price_asc' | 'price_desc' | 'name_asc' | 'name_desc' | 'newest' — допустимі алгоритми сортування каталогу.")

    add_paragraph(doc, "Типи об'єднання (Union Types) та техніка дискримінованих об'єднань (Discriminated Unions) використані для моделювання поліморфних структур даних:")

    add_paragraph(doc, "1. Дискриміноване об'єднання блоків динамічної сторінки (Page Builder). Для кожного типу блоку створено окремий інтерфейс (HeroBlockData, ProductGridBlockData, FeaturesBlockData, BannerBlockData, CtaBlockData, ModelViewerBlockData, CategoryCardsBlockData), що містить обов'язкове спільне дискримінаційне поле blockType із відповідним літеральним типом ('hero', 'productGrid' тощо). Загальний тип блоку визначено як:")

    add_paragraph(doc, "export type PageBlockData =\n"
                       "  | HeroBlockData\n"
                       "  | ProductGridBlockData\n"
                       "  | FeaturesBlockData\n"
                       "  | BannerBlockData\n"
                       "  | CtaBlockData\n"
                       "  | ModelViewerBlockData\n"
                       "  | CategoryCardsBlockData;")

    add_paragraph(doc, "Завдяки цій конструкції під час аналізу block.blockType усередині оператора switch компілятор TypeScript автоматично здійснює звуження типу (type narrowing), надаючи розробнику підказки щодо унікальних властивостей конкретного блоку (наприклад, доступ до масиву items для блоку features або поля filterType для блоку productGrid) та унеможливлюючи звернення до неіснуючих властивостей під час виконання.")

    add_paragraph(doc, "2. Узагальнене дискриміноване об'єднання станів інтерфейсу UIState<T>:\n"
                       "export type UIState<T> =\n"
                       "  | { status: 'idle' }\n"
                       "  | { status: 'loading' }\n"
                       "  | { status: 'success'; data: T }\n"
                       "  | { status: 'empty'; message?: string }\n"
                       "  | { status: 'error'; error: string; code?: string };\n"
                       "Такий тип гарантує, що поле data є доступним виключно тоді, коли значення status дорівнює 'success', а поле error — коли status дорівнює 'error'.")

    add_paragraph(doc, "3. Уніфікована обгортка серверних відповідей ApiResponse<T>:\n"
                       "export type ApiResponse<T> =\n"
                       "  | { success: true; data: T; message?: string; timestamp: string }\n"
                       "  | { success: false; error: string; code?: string; details?: Record<string, string[]>; timestamp: string };")

    add_placeholder(doc, "[ВСТАВИТИ ЗОБРАЖЕННЯ ЛІСТИНГУ 2.7: Система типів: Discriminated Unions та стани UI (cms/src/types/shop.ts)]")

    add_paragraph(doc, "Типи перетину (Intersection Types, оператор &) застосовано для побудови гнучких похідних структур даних без дублювання базових інтерфейсів:")

    add_paragraph(doc, "• Товар із розгорнутою сутністю категорії: за замовчуванням поле category у моделі Product містить ID або зв'язок. Тип ProductWithCategory = Product & { categoryDetails: Category } гарантує наявність повноцінного об'єкта категорії після виконання ORM-вибірки;\n"
                       "• Позиція кошика покупця: створюється на основі перетину ідентифікатора, кількості та безпечного знімка товару за допомогою утилітарного типу Pick:\n"
                       "  export type CartItemSnapshot = Pick<Product, 'name' | 'slug' | 'sku' | 'price' | 'imageUrl' | 'stock'>;\n"
                       "  export type CartItem = { productId: ID; quantity: number } & CartItemSnapshot;\n"
                       "• Елемент кошика з повними розрахунками вартості для інтерфейсу:\n"
                       "  export type CartItemWithDetails = CartItem & { product: Product; subtotal: number; isAvailable: boolean; hasEnoughStock: boolean };\n"
                       "• Розширений об'єкт замовлення з форматованими полями для перегляду:\n"
                       "  export type EnrichedOrder = Order & { formattedDate: string; itemsCount: number; isShipped: boolean };\n"
                       "• Сторінка з блоками та міткою оновлення: export type PageWithBlocks = PageLayout & { lastUpdatedFormatted?: string }.")

    add_paragraph(doc, "У модулі cms/src/lib/shopUtils.ts реалізовано строгу типізацію сигнатур функцій вищого порядку, предикатів та розрахункових операцій:")
    add_paragraph(doc, "1. Сигнатура валідатора кошика: export type CartValidatorFn = (items: Array<{ productId: ID; quantity: number }>, productsPool: Product[]) => CartValidationResult. Функція validateCart перевіряє активність кожної позиції, порівнює запитувану кількість із фактичним значенням product.stock і формує типізований результат CartValidationResult із переліком знайдених помилок.\n"
                       "2. Сигнатура калькулятора ціни: export type PriceCalculatorFn = (items: CartItemWithDetails[], deliveryService: DeliveryService) => PriceCalculationResult. Функція calculateOrderPricing підсумовує вартість товарів, призначає безкоштовну доставку при сумі замовлення від 1500 грн або застосовує тариф перевізника та повертає суворий об'єкт PriceCalculationResult.\n"
                       "3. Сигнатура фільтрації каталогу: export type FilterProductsFn = (products: Product[], options: ProductFilterOptions) => Product[]. Функція filterProducts виконує безпечний відбір за комплексним об'єктом опцій ProductFilterOptions, включаючи пошук за кількома текстовими полями одночасно, перевірку діапазону цін та сортування.")

    add_placeholder(doc, "[ВСТАВИТИ ЗОБРАЖЕННЯ ЛІСТИНГУ 2.8: Типізовані сервісні функції магазину (cms/src/lib/shopUtils.ts)]")

    add_paragraph(doc, "Типізація даних, отриманих через API, реалізується у двох вимірах. З боку бекенду Payload CMS використовує команду npm run generate:types, яка компілює схеми колекцій у монолітний файл payload-types.ts, що точно відображає структуру документів у базі даних PostgreSQL. З боку клієнтської вітрини застосовуються спеціалізовані типи запитів і відповідей (Data Transfer Objects, DTO), зокрема CheckoutApiRequest (контракт передачі даних покупця, адреси та масиву товарів) та CheckoutApiResponseData (контракт підтвердження замовлення з номером orderNumber та статусом).")

    # -------------------------------------------------------------------------
    # 2.4. Деплой та контрольний приклад
    # -------------------------------------------------------------------------
    add_heading_2(doc, "2.4. Деплой та контрольний приклад: розгортання проєкту, оформлення репозиторію (README.md), контрольний приклад із скріншотами (редактор в CMS, відвідувач на сайті, обробка помилок).")

    add_paragraph(doc, "Фінальним етапом створення Fullstack-застосунку GromBricks стало налаштування промислового середовища розгортання (Production Deployment), оформлення інженерної документації у файлі README.md та проведення комплексного тестування контрольного прикладу.")

    add_paragraph(doc, "Хмарна інфраструктура проєкту базується на тандемі платформ Neon та Vercel:")
    add_paragraph(doc, "• Хмарна база даних: у хмарному сервісі Neon створено екземпляр PostgreSQL 16. Отриманий рядок з'єднання містить параметри захищеного підключення sslmode=require та кінцеву точку пулера з'єднань PgBouncer, що гарантує миттєву обробку запитів під час пікових навантажень без вичерпання пулу сокетів бази даних.\n"
                       "• Розгортання у хмарі Vercel: з огляду на архітектуру монорепозиторію у панелі Vercel створено два пов'язаних проєкти з єдиного вихідного репозиторію GitHub:\n"
                       "  1. Проєкт бекенду (GromBricks CMS): коренева папка (Root Directory) вказана як cms. Змінні середовища: DATABASE_URL (адреса бази Neon), PAYLOAD_SECRET (секретний ключ хешування) та NEXT_PUBLIC_SERVER_URL (власний домен бекенду);\n"
                       "  2. Проєкт фронтенду (GromBricks Storefront): коренева папка (Root Directory) — frontend, команда збірки (Build Command) — npm run build, каталог виходу (Output Directory) — dist. Головна змінна середовища — VITE_API_BASE, яка вказує на публічний домен розгорнутого бекенду CMS.\n"
                       "• Подолання обмежень ефемерної файлової системи: використання зовнішньої СУБД Neon PostgreSQL у поєднанні з оптимізацією посилань на зображення дозволило повністю усунути ризик втрати даних, характерний для традиційних файлових сховищ на Vercel.")

    add_table_title(doc, "Таблиця 2.4 - Конфігураційні змінні оточення та параметри розгортання")
    add_table(doc,
              ["Застосунок у Vercel", "Параметр / Змінна оточення", "Значення / Конфігурація", "Призначення"],
              [
                  ["grombricks-cms (Backend)", "Root Directory", "cms", "Вказує Vercel на кореневу папку сервера Next.js."],
                  ["grombricks-cms (Backend)", "DATABASE_URL", "postgres://user:pass@ep-xyz.neon.tech/grombricks_db?sslmode=require", "Підключення до хмарної бази даних Neon PostgreSQL."],
                  ["grombricks-cms (Backend)", "PAYLOAD_SECRET", "32-символьний криптографічний хеш", "Захист JWT-токенів автентифікації користувачів."],
                  ["grombricks-cms (Backend)", "NEXT_PUBLIC_SERVER_URL", "https://grombricks-cms.vercel.app", "Публічна адреса адмін-панелі та API."],
                  ["grombricks-front (Frontend)", "Root Directory / Output", "frontend / dist", "Вказує шлях до вітрини Vite та каталогу збірки."],
                  ["grombricks-front (Frontend)", "VITE_API_BASE", "https://grombricks-cms.vercel.app", "Базова адреса для крос-доменних fetch-запитів до API."]
              ],
              col_widths=[1.6, 1.8, 2.0, 1.6])

    add_paragraph(doc, "Файл README.md у корені репозиторію оформлено за сучасними стандартами технічної документації. Він містить короткий опис місії проєкту GromBricks, детальну візуальну діаграму структури монорепозиторію, перелік мінімальних системних вимог (Node.js 20.18+, обліковий запис Neon), покрокову інструкцію з локального клонування та запуску проєкту, опис доступних команд npm-скриптів, зведену таблицю всіх кастомних API-ендпоінтів та інструкцію із розгортання на платформі Vercel.")

    add_paragraph(doc, "Для всебічної перевірки працездатності розробленого Fullstack-застосунку проведено тестування трьох взаємопов'язаних сценаріїв контрольного прикладу:")

    add_subheading(doc, "Сценарій 1. Робота адміністратора та редактора контенту в системі Payload CMS")
    add_paragraph(doc, "1. Адміністратор переходить за адресою http://localhost:3000/admin (або на хмарний домен Vercel) та успішно проходить авторизацію з обліковими даними адміністратора. На головній сторінці панелі відображається перелік доступних колекцій: Products, Categories, Orders, Pages, Media, Users.")
    add_placeholder(doc, "[ВСТАВИТИ РИСУНОК 2.2: Панель керування Payload CMS: список товарів колекції Products]")

    add_paragraph(doc, "2. У розділі Products адміністратор відкриває картку товару «Дарт Вейдер», редагує артикул SKU, встановлює статус рідкісності isRare у значення true та змінює складський залишок stock з 5 до 10 одиниць. Зміни миттєво зберігаються у базі даних Neon PostgreSQL.")
    add_paragraph(doc, "3. У розділі Pages адміністратор відкриває документ сторінки «home» та переглядає список блоків у Page Builder. Змінює заголовок у блоці Hero на «Ексклюзивні мініфігурки LEGO 2026», переставляє блок переваг Features вище блоку CTA та зберігає оновлену сторінку.")
    add_placeholder(doc, "[ВСТАВИТИ РИСУНОК 2.3: Візуальний конструктор блоків сторінки (Page Builder) у CMS]")

    add_paragraph(doc, "4. У глобальних налаштуваннях Globals -> HeaderSettings адміністратор коригує номер контактного телефону підтримки та адресу офіційного Telegram-каналу.")
    add_paragraph(doc, "5. У колекції Orders адміністратор знаходить щойно створене замовлення клієнта зі статусом pending, перевіряє розгорнутий список куплених деталей та змінює статус на processing («В обробці / комплектується»), фіксуючи зміну в системі.")
    add_placeholder(doc, "[ВСТАВИТИ РИСУНОК 2.4: Детальний перегляд замовлення та зміна статусу в колекції Orders]")

    add_subheading(doc, "Сценарій 2. Сценарій взаємодії покупця з вітриною інтернет-магазину")
    add_paragraph(doc, "1. Відвідувач відкриває головну сторінку вітрини index.html. Застосунок надсилає паралельні запити до /api/common та /api/page?path=/, після чого модульний шаблонізатор renderBlocks відображає оновлений Hero-банер, блок переваг та ініціалізує 3D-вітрину Three.js. Відвідувач взаємодіє з 3D-сценою, обертаючи деталізовану модель ніндзя Cole за допомогою миші.")
    add_placeholder(doc, "[ВСТАВИТИ РИСУНОК 2.5: Головна сторінка магазину GromBricks з інтерактивною 3D-вітриною]")

    add_paragraph(doc, "2. Відвідувач натискає кнопку «Перейти до каталогу» та опиняється на сторінці catalog.html. У полі пошуку вводить запит «Rex» та обирає чекбокс «Рідкісні ексклюзиви». Застосунок динамічно звертається до ендпоінту /api/products/search?q=Rex&onlyRare=true і миттєво оновлює товарну сітку, демонструючи колекційну фігурку Клон-капітана Рекса вартістю 1450 грн.")
    add_placeholder(doc, "[ВСТАВИТИ РИСУНОК 2.6: Каталог товарів із живим пошуком та фільтрацією за параметрами]")

    add_paragraph(doc, "3. Покупець натискає кнопку «В кошик». Скрипт cart.ts додає знімок товару до сховища localStorage, відтворює анімацію та оновлює числовий лічильник кошика у верхній панелі навігації з 0 на 1.")
    add_paragraph(doc, "4. Користувач переходить на сторінку cart.html, де відображається таблиця куплених товарів. За допомогою кнопки «+» кількість збільшується до 2 одиниць, при цьому проміжна сума автоматично перераховується з 1450 грн до 2900 грн.")
    add_placeholder(doc, "[ВСТАВИТИ РИСУНОК 2.7: Сторінка кошика покупця з розрахунком вартості]")

    add_paragraph(doc, "5. Натиснувши «Оформити замовлення», клієнт переходить на сторінку checkout.html, де заповнює форму: ім'я та прізвище («Олександр Мельник»), телефон (+380671112233), місто доставки («Львів»), номер відділення Нової Пошти («Відділення №15») та обирає спосіб оплати «Оплата при отриманні».")
    add_paragraph(doc, "6. Після натискання кнопки «Підтвердити замовлення» фронтенд надсилає POST-запит до ендпоінту /api/orders/checkout. Сервер успішно валідує наявність товару, списує 2 одиниці зі складу в базі даних Neon та повертає статус 201 із номером GB-2026-4819. Фронтенд очищує кошик і виводить екран подяки з номером замовлення.")
    add_placeholder(doc, "[ВСТАВИТИ РИСУНОК 2.8: Підтвердження успішно створеного замовлення]")

    add_subheading(doc, "Сценарій 3. Обробка граничних станів та помилок (Error Handling)")
    add_paragraph(doc, "1. Стан завантаження: при переході між розділами каталогу з імітацією повільного з'єднання (Slow 3G) користувачеві плавно транслюється анімаційний спінер .lego-spinner, що унеможливлює повторні хаотичні кліки.")
    add_paragraph(doc, "2. Стан порожньої видачі: під час пошуку за неіснуючим артикулом «NON-EXISTENT-SKU» сітка каталогу відображає блок .empty-container із повідомленням «Нічого не знайдено» та рекомендацією перевірити параметри запиту.")
    add_placeholder(doc, "[ВСТАВИТИ РИСУНОК 2.9: Стан інтерфейсу при відсутності результатів пошуку в каталозі]")

    add_paragraph(doc, "3. Клієнтська валідація форми: при спробі надіслати порожню форму замовлення на сторінці checkout.html браузер і скрипт checkoutPage.ts блокують відправку, виділяючи обов'язкові незаповнені поля.")
    add_paragraph(doc, "4. Серверний контроль складських залишків: при спробі замовити кількість товару, яка перевищує доступний залишок stock у базі даних (наприклад, 100 одиниць рідкісного шолома Мандалорця при залишку 10 шт.), ендпоінт /api/orders/checkout відхиляє операцію з кодом 400 Bad Request і повертає зрозумілу помилку. Вітрина відображає червоний банер .checkout-error-banner, захищаючи склад від перепродажу.")
    add_placeholder(doc, "[ВСТАВИТИ РИСУНОК 2.10: Повідомлення про помилку валідації залишків на складі при оформленні]")

    add_paragraph(doc, "5. Відмовостійкість вітрини (Fallback Mode): у разі штучного вимкнення сервера бекенду фронтенд-застосунок не припиняє виконання через необроблені проміси (Unhandled Promise Rejection). Завдяки сервісним заглушкам у модулі api.ts користувачеві відображається повноцінна статична головна сторінка з контактними даними та локальним кошиком.")

    # -------------------------------------------------------------------------
    # Висновок до розділу 2
    # -------------------------------------------------------------------------
    add_subheading(doc, "Висновок до розділу 2")
    add_paragraph(doc, "У другому розділі звіту спроєктовано, розроблено та протестовано повнофункціональний Fullstack-застосунок інтернет-магазину GromBricks на основі сучасної архітектури монорепозиторію з використанням Headless Payload CMS 3.x, хмарної СУБД Neon PostgreSQL, інструменту збірки Vite та строгої статичної типізації мовою TypeScript.")

    add_paragraph(doc, "На рівні бекенду сконфігуровано набір предметних колекцій (Products, Categories, Orders, Pages, Media, Users) та глобальних об'єктів (HeaderSettings, FooterSettings). Для подолання надмірної вкладеності стандартного REST API розроблено 6 кастомних API-ендпоінтів, що оптимізують передачу даних для вітрини та гарантують транзакційний контроль складських залишків. Проблему ефемерної файлової системи безсерверного хостингу Vercel успішно вирішено шляхом підключення хмарної реляційної бази даних Neon PostgreSQL із захищеним шифруванням SSL.")

    add_paragraph(doc, "Клієнтська частина вітрини реалізована за стандартом MPA на базі чистого Vanilla TypeScript та семантичної верстки HTML5/CSS3 без важких сторонніх фреймворків. Впроваджено гнучкий шаблонізатор сторінки (renderBlocks), інтерактивну 3D-вітрину Three.js з оглядом моделей LEGO, адаптивну адаптацію під мобільні екрани та надійну систему опрацювання життєвого циклу UI (завантаження, готовність, порожній стан, помилка). Повна статична типізація (Type Aliases, Discriminated Unions, Intersection Types) забезпечила надійність кодової бази та унеможливила виникнення помилок невідповідності типів на етапі виконання. Результати контрольного прикладу підтвердили повну відповідність розробленої системи усім поставленим практичним вимогам.")

if __name__ == '__main__':
    src_docx = r'C:\Users\vitai\.gemini\antigravity\brain\85df3c14-278b-43f5-ae1e-2ffdec46da01\.user_uploaded\media_1793085605805_620abf0d.docx'
    dst_root = r'd:\PracticeReport_3_TS\Звіт_з_практики_Розділ_1_і_2.docx'
    dst_docs = r'd:\PracticeReport_3_TS\docs\Звіт_з_практики_Розділ_1_і_2.docx'

    print(f"Loading original docx: {src_docx}")
    doc = docx.Document(src_docx)
    initial_paras = len(doc.paragraphs)
    print(f"Initial paragraph count: {initial_paras}")

    print("Building Section 2...")
    build_section_2(doc)
    final_paras = len(doc.paragraphs)
    print(f"Added {final_paras - initial_paras} paragraphs.")

    os.makedirs(os.path.dirname(dst_docs), exist_ok=True)
    doc.save(dst_root)
    print(f"Saved to: {dst_root}")
    doc.save(dst_docs)
    print(f"Saved to: {dst_docs}")
    print("Done!")
